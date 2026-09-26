# -*- coding: utf-8 -*-
"""Bộ đọc Office — EIDE-ING-43 §4.2. Bước ING-B.

Vì sao đây là bước đáng làm sớm: trong môi trường làm việc thật, spec nội bộ và bảng đo
là `.docx`/`.xlsx` nhiều hơn là PDF. Chính yêu cầu nâng cấp này cũng được gửi bằng
`.docx`. Một sản phẩm đọc được datasheet nhà sản xuất nhưng không đọc được tệp của
chính người dùng thì mới làm được nửa việc.

## Trích dẫn: đây là phần khó, không phải phần đọc chữ

Đọc chữ ra khỏi `.docx` là việc của một thư viện. Việc của bộ này là trả lời được câu
**"số này ở đâu trong tệp?"** theo cách người mở tệp ra tìm thấy:

| Loại | Đơn vị trích dẫn | Ví dụ |
|---|---|---|
| `.docx` | đường tiêu đề + số bảng/đoạn | `3.2 Electrical > Bảng 4` |
| `.xlsx` | sheet và ô | `Bảng đo!B7` |
| `.pptx` | số slide | `slide 12` |

`.docx` **không có số trang cố định** — nó phụ thuộc phông chữ, khổ giấy, máy in. Nên
trích dẫn theo trang cho Word là một lời hứa sai; ai mở trên máy khác sẽ thấy số khác.
Khi người thật sự cần số trang (để in, để đối chiếu với bản giấy), ta sinh một bản PDF
phái sinh **cùng `doc_id`** và nói rõ đó là bản phái sinh.

## Công thức trong Excel là NGUỒN CỦA SỐ

`openpyxl` đọc được cả công thức lẫn giá trị. Một ô ghi `=AVERAGE(B2:B9)` nói với người
rà soát nhiều hơn con số `3.27` — nó cho biết số đó **đến từ đâu**. Nên công thức được
giữ làm provenance, không bị vứt.
"""

from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Any

from .docs import _MAU_TIEM_LENH, TaiLieu, Trang

# Tiêu đề cột báo hiệu một bảng THÔNG SỐ (khác bảng trình bày). ING-43 §5.1 bước 1.
_TU_DIEN_COT = {
    "parameter", "symbol", "min", "typ", "typical", "max", "maximum", "minimum",
    "unit", "units", "conditions", "condition", "value", "rating",
    "thông số", "ký hiệu", "đơn vị", "giá trị", "điều kiện", "nhỏ nhất", "lớn nhất",
}

# Bảng BẢN ĐỒ CHÂN là một loại bảng khác hẳn bảng thông số, và nó là loại bảng quan trọng
# nhất trong một tài liệu BÀN GIAO PHẦN CỨNG: nó nói chân nào nối đi đâu.
#
# Đo được trên tài liệu thật (MOBILUCK, bảng 12 — 23 chân): từ điển cũ chỉ có tên cột của
# datasheet điện (parameter/min/typ/max/unit), nên bảng `Chân | Hướng | Net · khối | Chức
# năng` KHÔNG được nhận là bảng — mọi hàng bị đọc như một dòng chữ, và cả bản đồ chân biến
# mất khỏi phần trích xuất. Tác tử sau đó chỉ trích được 7 Fact, không Fact nào là chân.
_TU_DIEN_COT_CHAN = {
    "chân", "chan", "pin", "chân số", "số chân", "pin number",
    "hướng", "huong", "direction", "i/o", "io", "dir",
    "net", "tín hiệu", "tin hieu", "signal", "net · khối", "net/khối", "kết nối",
    "chức năng", "chuc nang", "function", "mô tả", "ghi chú", "cổng", "port",
}

MAX_O_MOI_SHEET = 20_000        # bảng đo lớn thì cắt và NÓI RA, không đọc vô hạn
MAX_SLIDE = 500


def _toa_do(hang: tuple[Any, ...]) -> str:
    """Toạ độ ô đầu tiên của một hàng, chịu được `EmptyCell` của chế độ read-only."""
    for o in hang:
        tđ = getattr(o, "coordinate", None)
        if tđ:
            return str(tđ)
    r = next((getattr(o, "row", None) for o in hang if getattr(o, "row", None)), "?")
    return f"A{r}"


def _bam(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _canh_bao_tiem_lenh(dv: list[Trang]) -> list[str]:
    ra: list[str] = []
    for t in dv:
        for mau in _MAU_TIEM_LENH:
            m = mau.search(t.chu)
            if m:
                ra.append(f"{t.trich_dan}: “{m.group(0)[:60]}”")
                break
    return ra


def la_bang_thong_so(hang_dau: list[str]) -> bool:
    """Bảng này có phải bảng thông số không — quyết định bằng TIÊU ĐỀ CỘT.

    §5.1: tiêu đề khớp từ điển thì đọc giá trị bằng mã; không khớp thì mô hình chỉ được
    ánh xạ tiêu đề cột, **không đọc giá trị**. Ranh giới đó là lý do số trích ra tin
    được: máy đọc số, mô hình chỉ đặt tên cột.
    """
    thap = {str(c).strip().lower() for c in hang_dau if str(c).strip()}
    return len(thap & _TU_DIEN_COT) >= 2


def la_bang_chan(hang_dau: list[str]) -> bool:
    """Bảng này có phải BẢN ĐỒ CHÂN không — cũng quyết định bằng tiêu đề cột.

    Đòi hai cột khớp, và một trong hai phải là cột "chân": một bảng chỉ có "Chức năng" và
    "Ghi chú" là bảng mô tả, không phải bản đồ chân.
    """
    thap = {str(c).strip().lower() for c in hang_dau if str(c).strip()}
    co_chan = any(c.split("·")[0].strip() in ("chân", "chan", "pin", "chân số", "số chân",
                                              "pin number")
                  for c in thap)
    return co_chan and len(thap & _TU_DIEN_COT_CHAN) >= 2


def loai_bang(hang_dau: list[str]) -> str:
    """`"thong_so"` · `"chan"` · `""` (bảng trình bày)."""
    if la_bang_thong_so(hang_dau):
        return "thong_so"
    return "chan" if la_bang_chan(hang_dau) else ""


# =========================================================================== .docx
def doc_docx(path: Path, *, doc_id: str, phien_ban: str = "",
             nha_phat_hanh: str = "") -> TaiLieu:
    """Đoạn theo cây tiêu đề; bảng gắn với tiêu đề gần nhất."""
    import docx                                              # python-docx

    d = docx.Document(str(path))
    duong: list[str] = []          # cây tiêu đề hiện tại: H1 > H2 > H3
    dv: list[Trang] = []
    so = 0
    so_bang = 0

    def nhan_hien_tai(them: str = "") -> str:
        goc = " > ".join(duong) if duong else "(đầu tài liệu)"
        return f"{goc} > {them}" if them else goc

    # python-docx không cho duyệt xen kẽ đoạn và bảng qua API cấp cao, nên đi theo thứ
    # tự phần tử XML của thân tài liệu — thứ tự đó mới là thứ tự người đọc nhìn thấy.
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for phan_tu in d.element.body.iterchildren():
        if phan_tu.tag.endswith("}p"):
            p = Paragraph(phan_tu, d)
            chu = p.text.strip()
            if not chu:
                continue
            # Đoạn có thể KHÔNG có style (`p.style` là None) — gặp ngay trên tệp
            # thật của chủ sản phẩm. Một bộ đọc tài liệu không được phép nổ vì
            # một đoạn thiếu định dạng.
            ten_style = getattr(getattr(p, "style", None), "name", "") or ""
            m = re.match(r"Heading (\d)", ten_style)
            if m:
                muc = int(m.group(1))
                duong[:] = duong[: muc - 1]
                duong.append(chu)
                continue
            so += 1
            dv.append(Trang(so, chu, nhan=nhan_hien_tai(f"đoạn {so}")))
        elif phan_tu.tag.endswith("}tbl"):
            t = Table(phan_tu, d)
            so_bang += 1
            hang = [[o.text.strip() for o in r.cells] for r in t.rows]
            if not hang:
                continue
            lb = loai_bang(hang[0])
            la_ts = bool(lb)
            for i, r in enumerate(hang):
                if not any(r):
                    continue
                so += 1
                # Giữ Ô và TIÊU ĐỀ CỘT, không chỉ giữ chuỗi đã nối.
                #
                # Trong datasheet, đơn vị nằm ở **cột riêng**: "VDD max | 2.7 | 5.5 | V".
                # Nối thành một dòng rồi tìm "số kèm đơn vị" thì không bao giờ khớp, vì
                # `5.5` và `V` cách nhau một dấu gạch. Bảng phải được đọc như bảng.
                dv.append(Trang(
                    so, " | ".join(r),
                    nhan=nhan_hien_tai(f"Bảng {so_bang}, dòng {i + 1}"
                                       + ("" if la_ts else " (bảng trình bày)")),
                    o=list(r), cot=list(hang[0]) if la_ts and i > 0 else [],
                    loai_bang=lb if i > 0 else ""))

    props = d.core_properties
    return TaiLieu(
        doc_id=doc_id, ten=path.name, duong_dan=str(path), hash=_bam(path),
        so_trang=len(dv), trang=dv, loai="docx", don_vi_trich_dan="mục",
        phien_ban=phien_ban or (props.revision and str(props.revision) or ""),
        nha_phat_hanh=nha_phat_hanh or (props.author or ""),
        canh_bao_tiem_lenh=_canh_bao_tiem_lenh(dv))


# =========================================================================== .xlsx
def doc_xlsx(path: Path, *, doc_id: str, phien_ban: str = "",
             nha_phat_hanh: str = "") -> TaiLieu:
    """Mỗi sheet một vùng dữ liệu; mỗi hàng một đơn vị trích dẫn `Sheet!ô`."""
    import openpyxl

    # Đọc HAI lần: một lần lấy giá trị đã tính, một lần lấy công thức, rồi ghép theo
    # từng hàng. Công thức là nguồn của số (§4.2) — vứt nó đi là vứt mất câu trả lời
    # "số này ở đâu ra".
    #
    # Có một trường hợp dễ mất dữ liệu trong im lặng: ô chỉ chứa công thức mà tệp CHƯA
    # TỪNG được Excel mở thì không có giá trị đã tính, `data_only=True` trả về None, và
    # cả hàng bị coi là rỗng rồi biến mất. Nên chỗ nào không có giá trị thì lấy công
    # thức làm nội dung — thà hiện `=AVERAGE(B2:B3)` còn hơn không hiện gì.
    wb_gt = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    wb_ct = openpyxl.load_workbook(str(path), data_only=False, read_only=True)

    dv: list[Trang] = []
    so = 0
    cat_bot = 0
    try:
        for ten_sheet in wb_gt.sheetnames:
            sh = wb_gt[ten_sheet]
            sh_ct = wb_ct[ten_sheet] if ten_sheet in wb_ct.sheetnames else None
            hang_ct = sh_ct.iter_rows(values_only=True) if sh_ct is not None else None
            n_o = 0
            for hang in sh.iter_rows(values_only=False):
                ct_hang = next(hang_ct, ()) if hang_ct is not None else ()
                phan: list[str] = []
                for i, o in enumerate(hang):
                    v = o.value
                    ct = ct_hang[i] if i < len(ct_hang) else None
                    la_ct = isinstance(ct, str) and ct.startswith("=")
                    if v is None or not str(v).strip():
                        if la_ct:
                            phan.append(str(ct))     # chưa tính — công thức LÀ nội dung
                        continue
                    phan.append(f"{v} [{ct}]" if la_ct else str(v))
                if not phan:
                    continue
                n_o += len(hang)
                if n_o > MAX_O_MOI_SHEET:
                    cat_bot += 1
                    break
                so += 1
                # Ô đầu hàng có thể là `EmptyCell` (read-only) và không có `coordinate`
                # — lấy ô đầu tiên CÓ toạ độ. Đây là chỗ trích dẫn được dựng, nên nó
                # không được phép hỏng vì một ô trống ở đầu bảng.
                dv.append(Trang(so, " | ".join(phan),
                                nhan=f"{ten_sheet}!{_toa_do(hang)}"))
    finally:
        wb_gt.close()
        wb_ct.close()

    canh = _canh_bao_tiem_lenh(dv)
    if cat_bot:
        canh.append(f"{cat_bot} sheet bị cắt ở {MAX_O_MOI_SHEET:,} ô — bảng quá lớn.")
    return TaiLieu(
        doc_id=doc_id, ten=path.name, duong_dan=str(path), hash=_bam(path),
        so_trang=len(dv), trang=dv, loai="xlsx", don_vi_trich_dan="ô",
        phien_ban=phien_ban, nha_phat_hanh=nha_phat_hanh,
        canh_bao_tiem_lenh=canh)


# =========================================================================== .pptx
def doc_pptx(path: Path, *, doc_id: str, phien_ban: str = "",
             nha_phat_hanh: str = "") -> TaiLieu:
    """Slide: tiêu đề + thân + ghi chú. Mức MỘT PHẦN — không trích Fact tự động."""
    from pptx import Presentation

    pr = Presentation(str(path))
    dv: list[Trang] = []
    for i, slide in enumerate(pr.slides, 1):
        if i > MAX_SLIDE:
            break
        phan: list[str] = []
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                phan.append(sh.text_frame.text.strip())
            if getattr(sh, "has_table", False):
                for r in sh.table.rows:
                    phan.append(" | ".join(o.text.strip() for o in r.cells))
        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                phan.append("[ghi chú] "
                            + slide.notes_slide.notes_text_frame.text.strip())
        except Exception:                                    # noqa: BLE001
            pass
        if phan:
            dv.append(Trang(i, "\n".join(phan), nhan=f"slide {i}"))

    return TaiLieu(
        doc_id=doc_id, ten=path.name, duong_dan=str(path), hash=_bam(path),
        so_trang=len(dv), trang=dv, loai="pptx", don_vi_trich_dan="slide",
        phien_ban=phien_ban, nha_phat_hanh=nha_phat_hanh,
        canh_bao_tiem_lenh=_canh_bao_tiem_lenh(dv))


# =========================================================================== chuyển đổi
def co_libreoffice() -> str | None:
    """Tìm LibreOffice. Không có thì trả None — và chỗ gọi phải NÓI RA, không nổ."""
    import shutil

    for ten in ("soffice", "libreoffice"):
        d = shutil.which(ten)
        if d:
            return d
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


def chuyen_doi(path: Path, sang: str, thu_muc_ra: Path) -> tuple[Path | None, str]:
    """LibreOffice headless: `.doc`→`.docx`, `.docx`→`.pdf`…

    Trả `(đường dẫn kết quả, lý do nếu hỏng)`. Không ném ngoại lệ: thiếu LibreOffice là
    một tình huống bình thường trên máy người dùng, không phải một sự cố của sản phẩm.
    """
    import subprocess

    exe = co_libreoffice()
    if exe is None:
        return None, ("Máy chưa có LibreOffice. Định dạng Office đời cũ cần nó để đọc; "
                      "cách khác là lưu lại tệp ở định dạng mới (.docx/.xlsx/.pptx).")
    thu_muc_ra.mkdir(parents=True, exist_ok=True)
    try:
        r = subprocess.run(
            [exe, "--headless", "--convert-to", sang, "--outdir",
             str(thu_muc_ra), str(path)],
            capture_output=True, text=True, timeout=180)
    except subprocess.TimeoutExpired:
        return None, "LibreOffice chạy quá 180 giây — tệp có thể quá lớn hoặc hỏng."
    except OSError as e:
        return None, f"Không chạy được LibreOffice: {e}"

    ra = thu_muc_ra / (path.stem + "." + sang.split(":")[0])
    if not ra.exists():
        return None, (f"LibreOffice không tạo được tệp {sang}: "
                      f"{(r.stderr or r.stdout or '').strip()[:200]}")
    return ra, ""


# =========================================================================== điều phối
_BO_DOC = {"docx": doc_docx, "xlsx": doc_xlsx, "pptx": doc_pptx}
_DOI_SANG = {".doc": ("docx", doc_docx), ".xls": ("xlsx", doc_xlsx),
             ".ppt": ("pptx", doc_pptx), ".odt": ("docx", doc_docx),
             ".ods": ("xlsx", doc_xlsx), ".odp": ("pptx", doc_pptx),
             ".rtf": ("docx", doc_docx)}


def nap_office(path: Path, *, loai: str, doc_id: str, thu_muc_tam: Path,
               phien_ban: str = "", nha_phat_hanh: str = "") -> tuple[TaiLieu | None, str]:
    """Nạp một tệp Office. Trả `(TaiLieu, lý do nếu không đọc được)`."""
    if loai in _BO_DOC:
        return _BO_DOC[loai](path, doc_id=doc_id, phien_ban=phien_ban,
                             nha_phat_hanh=nha_phat_hanh), ""

    if loai == "office_cu":
        duoi = path.suffix.lower()
        if duoi not in _DOI_SANG:
            return None, f"Chưa biết chuyển {duoi} sang định dạng nào."
        sang, doc_ham = _DOI_SANG[duoi]
        moi, loi = chuyen_doi(path, sang, thu_muc_tam)
        if moi is None:
            return None, loi
        tl = doc_ham(moi, doc_id=doc_id, phien_ban=phien_ban,
                     nha_phat_hanh=nha_phat_hanh)
        # Hash theo tệp GỐC — đó mới là thứ người dùng đưa vào và sẽ đối chiếu.
        tl.hash = _bam(path)
        tl.ten = path.name
        tl.duong_dan = str(path)
        tl.chuyen_doi_tu = duoi
        return tl, ""

    return None, f"Không có bộ đọc cho loại {loai}."


def lam_pdf_phai_sinh(path: Path, thu_muc_ra: Path) -> tuple[Path | None, str]:
    """`.docx` → PDF **phái sinh**, để có số trang khi người cần in hoặc đối chiếu.

    Gọi là "phái sinh" chứ không phải "bản gốc" vì nó là một cách trình bày của cùng
    một nội dung: số trang phụ thuộc phông chữ và khổ giấy, nên nó chỉ đúng cho bản PDF
    này. Trích dẫn chính vẫn là đường tiêu đề.
    """
    return chuyen_doi(path, "pdf", thu_muc_ra)
