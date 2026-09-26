#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Kiểm ING-C/D/E và MEM-D qua giao diện thật.

    python tools/thu_cuoi.py

Happy: chuẩn hoá `4R7`/`3V3`/dải · bảng Office → Fact có điều kiện · netlist + BOM đối
chiếu · `.ld` lệch datasheet · ngôn ngữ và gói OCR · C4 khẩn cấp · dọn theo tham chiếu ·
sáu chỉ số đo.

Unhappy (chín đường): `TBD` thành 0 · `100n` không ngữ cảnh · BOM thiếu cột mã · netlist
sai định dạng · tầng CẤU HÌNH làm vế so sánh · thiếu gói OCR · gc xoá blob còn tham
chiếu · đo lường báo đạt khi chưa đủ dữ liệu · C4 chạm message ghim.
"""

from __future__ import annotations

import json
import pathlib
import shutil
import subprocess
import sys
import time

REPO = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "src"))
sys.path.insert(0, str(REPO / "tools"))
sys.path.insert(0, str(REPO / "tests"))

from thu_giao_dien import Bo, GiaoDien        # noqa: E402

XANH, HET = "\033[92m", "\033[0m"
EX = {"summary": "nạp", "why": "để dùng số có nguồn", "sources": [], "diff_prev": "—",
      "next": "—", "confidence": "NGUOI"}

_NET = """(export (version D)
  (components
    (comp (ref R5) (value 4k7) (footprint R_0603))
    (comp (ref C4) (value 100nF) (footprint C_0402))
    (comp (ref U1) (value ATmega328P)))
  (nets
    (net (code 1) (name "+3V3") (node (ref U1) (pin 7)) (node (ref C4) (pin 1)))
    (net (code 2) (name "SDA") (node (ref U1) (pin 27)) (node (ref R5) (pin 1)))
    (net (code 3) (name "HO") (node (ref R5) (pin 2)))))
"""


def chay(du_an: pathlib.Path) -> int:
    b = Bo()
    g = GiaoDien(du_an)
    print("Mở app…")
    subprocess.run(["open", str(REPO / "ui/EIDEApp/EIDE.app")], check=True)
    g.san_sang(60)
    print(f"{XANH}Kênh kiểm thử giao diện đã mở{HET}\n")

    from eide.tools import build_registry
    reg = build_registry()
    ctx = _ctx(du_an)

    # ================================================================== ING-C
    b.phan("A · CHUẨN HOÁ ĐƠN VỊ VÀ KÝ HIỆU (ING12)")
    from eide.knowledge.chuan_hoa import chuan_hoa
    mau = [("4R7", 4.7, "Ω"), ("3V3", 3.3, "V"), ("100 nF", 100e-9, "F"),
           ("4,7kohm", 4700, "Ω"), ("20 mA", 0.02, "A")]
    ok = all(abs((chuan_hoa(c).gia_tri or 0) - v) < 1e-12 and chuan_hoa(c).don_vi == d
             for c, v, d in mau)
    b.kiem("Ký hiệu kỹ thuật và tiền tố SI đọc đúng", ok,
           "; ".join(f"{c}→{chuan_hoa(c).hien_vi()}" for c, _, _ in mau))

    g1 = chuan_hoa("2.7–5.5 V")
    b.kiem("Dải tách thành min/max",
           g1.la_dai and g1.vmin == 2.7 and g1.vmax == 5.5, g1.hien_vi())
    b.kiem("Hiện cho người Việt dùng dấu phẩy", chuan_hoa("3.3 V").hien_vi() == "3,3 V",
           chuan_hoa("3.3 V").hien_vi())
    b.kiem("M hoa và m thường khác nhau một tỉ lần",
           chuan_hoa("1 MHz").gia_tri == 1e6 and chuan_hoa("1 mHz").gia_tri == 1e-3,
           f"1 MHz={chuan_hoa('1 MHz').gia_tri:g} · 1 mHz={chuan_hoa('1 mHz').gia_tri:g}")

    b.phan("B · BẢNG OFFICE → FACT CÓ ĐIỀU KIỆN (ING09)")
    r = reg.run("doc.load", {"path": "spec.docx", "doc_id": "SPEC", "nguon": "noi_bo",
                             "explain": EX}, ctx)
    b.kiem("Nạp được spec nội bộ", r.ok, str(r.error)[:110] if not r.ok else
           f"{r.data['so_don_vi']} mục")
    r = reg.run("fact.extract", {"doc_id": "SPEC", "thuc_the": "chip:ATmega328P"}, ctx)
    b.kiem("Trích được Fact từ bảng", r.ok and r.data["so_ung_vien"] > 0,
           f"{r.data.get('so_ung_vien')} ứng viên" if r.ok else str(r.error)[:100])
    if r.ok and r.data["so_ung_vien"]:
        theo = {u["khoa"]: u for u in r.data["fact"]}
        b.kiem("min/typ/max chung một ô tách đúng",
               {"vdd.min", "vdd.max"} <= set(theo),
               ", ".join(sorted(theo)[:6]))
        co_dk = [u for u in r.data["fact"] if u.get("dieu_kien")]
        b.kiem("Điều kiện đi theo Fact", bool(co_dk),
               str(co_dk[0]["dieu_kien"]) if co_dk else "không có Fact nào mang điều kiện")

    b.phan("C · ĐỐI CHIẾU CHÉO NGUỒN (ING16)")
    ctx.store.put_fact({"fact_id": "f_cu", "subject": "chip:ATmega328P", "key": "vdd.max",
                        "value": 6.0, "unit": "V", "tier": "BAC", "origin": "extract",
                        "source": {"doc_id": "DS-cu"}, "explain": {}})
    ctx.store.put_fact({"fact_id": "f_er", "subject": "chip:ATmega328P", "key": "vdd.max",
                        "value": 5.5, "unit": "V", "tier": "BAC", "origin": "extract",
                        "source": {"doc_id": "ERRATA-01", "errata": True}, "explain": {}})
    r = reg.run("fact.cross_check", {}, ctx)
    b.kiem("Phát hiện hai nguồn nói hai số", r.ok and r.data["so_lech"] >= 1,
           f"{r.data.get('so_lech')} chỗ lệch" if r.ok else "")
    if r.ok and r.data["so_lech"]:
        d = r.data["lech"][0]
        b.kiem("Đề xuất errata, kèm lý do",
               d["de_xuat"] == "f_er" and "errata" in d["vi_sao"].lower(),
               f"{d['de_xuat']} · {d['vi_sao'][:70]}")
        b.kiem("KHÔNG tự chọn — bắt hỏi người",
               d["can_nguoi_chon"] and "Đừng tự chọn" in r.data["note_vi"],
               r.data["note_vi"][:100])

    # ================================================================== ING-D
    b.phan("D · EDA: NETLIST VÀ BOM (ING07, TC062)")
    r = reg.run("eda.netlist", {"path": "mach.net", "explain": EX}, ctx)
    b.kiem("Đọc netlist thành cấu trúc", r.ok and r.data["so_linh_kien"] == 3,
           f"{r.data.get('so_linh_kien')} linh kiện · {r.data.get('so_net')} net"
           if r.ok else str(r.error)[:100])
    b.kiem("Net chỉ nối một chân bị nêu ra",
           r.ok and r.data["net_mot_chan"] == ["HO"],
           str(r.data.get("net_mot_chan")) if r.ok else "")

    r = reg.run("eda.bom_check", {"bom": "bom.csv", "netlist": "netlist:mach.net"}, ctx)
    b.kiem("BOM thiếu linh kiện bị phát hiện",
           r.ok and r.data["thieu_trong_bom"] == ["C4"],
           str(r.data.get("thieu_trong_bom")) if r.ok else str(r.error)[:100])
    b.kiem("Giá trị khác nhau — loại tệ nhất — bị nêu đích danh",
           r.ok and r.data["gia_tri_khac"]
           and r.data["gia_tri_khac"][0]["ref"] == "R5",
           str(r.data.get("gia_tri_khac"))[:110] if r.ok else "")
    b.kiem("Nói rõ hậu quả: mạch hàn xong CHẠY nhưng sai",
           r.ok and any("CHẠY nhưng sai" in d for d in r.data["se_mat_vi"]),
           "; ".join(r.data.get("se_mat_vi", []))[:120] if r.ok else "")

    b.phan("Đ · TẦNG CẤU HÌNH VÀ LỆCH VỚI DATASHEET (ING08, ING11)")
    ctx.store.put_fact({"fact_id": "f_flash", "subject": "chip:ATmega328P",
                        "key": "flash.size", "value": 32768, "unit": "B", "tier": "BAC",
                        "origin": "extract", "source": {"doc_id": "DS", "page": 12},
                        "explain": {}})
    r = reg.run("config.load", {"path": "link.ld", "explain": EX}, ctx)
    b.kiem("Đọc linker script thành Fact CẤU HÌNH",
           r.ok and r.data["tang"] == "CAUHINH",
           f"{r.data.get('so_khoa')} khoá · tầng {r.data.get('tang')}" if r.ok else
           str(r.error)[:110])
    b.kiem(".ld khai 64 K mà chip có 32 K → PHÁT HIỆN",
           r.ok and any(l["muc"] == "vuot" for l in r.data["lech_voi_tai_lieu"]),
           r.data["lech_voi_tai_lieu"][0]["message_vi"][:130]
           if r.ok and r.data["lech_voi_tai_lieu"] else "không phát hiện")
    b.kiem("Nói rõ lỗi này không lộ ra lúc biên dịch",
           r.ok and any("không lộ ra lúc biên dịch" in l["message_vi"]
                        for l in r.data["lech_voi_tai_lieu"]),
           "")

    b.phan("E · NGÔN NGỮ VÀ GÓI OCR (ING11)")
    from eide.knowledge.ocr import nhan_ngon_ngu
    b.kiem("Nhận đúng ngôn ngữ và giải trình được",
           nhan_ngon_ngu("工作电压为3.3V").ma == "chi_sim"
           and nhan_ngon_ngu("Điện áp hoạt động").ma == "vie"
           and bool(nhan_ngon_ngu("Điện áp hoạt động").vi_sao),
           nhan_ngon_ngu("Điện áp hoạt động").vi_sao)
    from eide.knowledge.ocr import khoa_tu_bi_danh
    b.kiem("Bí danh đa ngữ ánh xạ đúng khoá chuẩn",
           khoa_tu_bi_danh("工作电压") == "vdd.range"
           and khoa_tu_bi_danh("điện áp hoạt động") == "vdd.range",
           "工作电压 → vdd.range")

    # ================================================================== MEM-D
    b.phan("G · MEM-D: C4, DỌN RÁC, ĐO LƯỜNG")
    from eide.memory.nen import c4
    ms = [{"role": "user", "text": "x", "_kind": "say"},
          {"role": "tool", "tool": "fs.read", "tool_call_id": "c1",
           "result": {"ok": True, "data": {"content": "nội dung " * 400}},
           "envelope": {"summary_line": "fs.read a.c",
                        "blob_ref": "blob:sha256:" + "1" * 64}}]
    bc = c4(ms)
    b.kiem("C4 bỏ kết quả công cụ nhưng giữ đường đọc lại",
           bc["stub"] == 1 and "blob.read" in ms[1]["result"]["_da_thu_gon"],
           f"giảm {bc['giam_phan_tram']:.0f} %")

    r = reg.run("memory.gc", {}, ctx)
    b.kiem("Dọn rác mặc định chỉ ĐỀ XUẤT", r.ok and r.data["thu"] is True,
           r.data.get("dong_vi", "")[:110] if r.ok else "")
    b.kiem("Nói rõ dọn theo tham chiếu, không theo tuổi",
           r.ok and "THAM CHIẾU, không theo tuổi" in r.data["note_vi"], "")

    r = reg.run("memory.metrics", {}, ctx)
    b.kiem("Sáu chỉ số đo có mục tiêu đọc được",
           r.ok and all(x["muc_tieu"] for x in r.data["danh_gia"]),
           "; ".join(f"{x['chi_so']}: {x['ket_qua']}" for x in r.data["danh_gia"][:2])
           if r.ok else "")

    b.phan("H · TÁC TỬ THẬT DÙNG CÁC CÔNG CỤ MỚI")
    g.go("Trong dự án có mach.net và bom.csv. Đối chiếu giúp mình xem BOM có khớp sơ đồ "
         "không, và nói rõ chỗ nào lệch thì hậu quả là gì.")
    a = g.doi_xong(400)
    loi = a["loi_tac_tu_cuoi"]
    b.kiem("Tác tử nêu được linh kiện thiếu và giá trị lệch",
           ("C4" in loi and "R5" in loi), loi[:170])

    # ================================================================== UNHAPPY
    b.phan("I · CHÍN ĐƯỜNG HỎNG")

    b.buoc("1. TBD không được thành 0")
    gg = chuan_hoa("TBD")
    b.kiem("Trả null có LÝ DO, không trả 0",
           not gg.co_so and "TBD" in gg.co, gg.co[:90])

    b.buoc("2. 100n không có ngữ cảnh")
    gg = chuan_hoa("100n")
    b.kiem("Không bịa đơn vị", gg.gia_tri is None and "không rõ đơn vị" in gg.co,
           gg.co[:90])

    b.buoc("3. BOM thiếu cột mã linh kiện")
    r = reg.run("eda.bom_check", {"bom": "bom-thieu-cot.csv",
                                  "netlist": "netlist:mach.net"}, ctx)
    b.kiem("Từ chối, không đoán cột nào là cột nào",
           not r.ok and "không đoán" in r.error.message_vi.lower(),
           r.error.message_vi[:110] if r.error else "lọt!")

    b.buoc("4. Netlist sai định dạng")
    r = reg.run("eda.netlist", {"path": "khong-phai.net", "explain": EX}, ctx)
    b.kiem("Nói đúng là không phải netlist KiCad",
           not r.ok and "netlist KiCad" in r.error.message_vi,
           r.error.message_vi[:110] if r.error else "lọt!")

    b.buoc("5. Tầng CẤU HÌNH làm vế so sánh giới hạn vật lý")
    ch = [f for f in ctx.store.query_facts(limit=200) if f["tier"] == "CAUHINH"]
    if ch:
        r = reg.run("fact.compare", {"luat": "ngan_sach_bo_nho", "fact_a": "f_flash",
                                     "fact_b": ch[0]["fact_id"]}, ctx)
        # Phải từ chối VÌ TẦNG, không vì tên luật lạ — nếu không, ô này đậu vì lý do sai.
        b.kiem("Từ chối kết luận — CẤU HÌNH không phải giới hạn vật lý",
               r.ok and r.data["chua_kiem_chung"]
               and "luật" not in r.data.get("giai_thich", ""),
               str(r.data.get("giai_thich"))[:130] if r.ok else "")

    b.buoc("6. Thiếu gói OCR cho ngôn ngữ tài liệu")
    from eide.knowledge.ocr import goi_da_cai, kiem_goi
    duoc, vi_sao = kiem_goi("vie")
    if "vie" not in goi_da_cai():
        b.kiem("Nói thẳng thiếu gói, không OCR bừa",
               not duoc and "nhìn như chữ mà sai" in vi_sao, vi_sao[:130])
        b.kiem("Chỉ luôn cách cài", "tesseract" in vi_sao, vi_sao[-90:])
    else:
        b.kiem("Máy có gói vie — OCR được", duoc, "có gói")

    b.buoc("7. gc không được xoá blob còn tham chiếu")
    import hashlib
    import os
    noi = b"bang chung cua snapshot"
    h = hashlib.sha256(noi).hexdigest()
    pb = ctx.config.paths.blobs / h[:2] / h
    pb.parent.mkdir(parents=True, exist_ok=True)
    pb.write_bytes(noi)
    t = time.time() - 400 * 86400
    os.utime(pb, (t, t))
    (ctx.config.paths.state_dir / "snapshots.jsonl").write_text(
        json.dumps({"id": "snap-01", "contents": {"store_export_hash": h}}) + "\n",
        "utf-8")
    from eide.memory import gc as _gc
    bcg = _gc(ctx.config.paths, thu=True)
    b.kiem("Blob 400 ngày nhưng còn snapshot trỏ tới thì GIỮ",
           h not in bcg.blob_se_xoa,
           f"{bcg.blob_co_tham_chieu} blob có tham chiếu")

    b.buoc("8. Đo lường khi chưa đủ dữ liệu")
    from eide.memory import do_luong
    from eide.protocol.ledger import Ledger
    so_moi = Ledger(du_an / ".eide" / "so-gia.jsonl")
    so_moi.append("turn.end", {"run_id": "r1", "cost": {"tokens": {"vao": 100}}})
    bang = do_luong(so_moi).dat_khong()
    b.kiem("Nói CHƯA ĐỦ, không nói đạt",
           all(x["ket_qua"] == "chưa đủ dữ liệu" for x in bang),
           "; ".join(x["ket_qua"] for x in bang))

    b.buoc("9. C4 chạm message ghim")
    ms2 = [{"role": "tool", "tool": "fs.read", "tool_call_id": "c1",
            "result": {"ok": True, "data": {"content": "x" * 900}},
            "envelope": {"summary_line": "fs.read a.c"}},
           {"role": "tool", "tool": "fs.read", "tool_call_id": "c2",
            "result": {"ok": True, "data": {"content": "y" * 900}},
            "envelope": {"summary_line": "fs.read b.c"}}]
    bc2 = c4(ms2, ghim={0})
    b.kiem("Ghim thì C4 cũng không chạm",
           not ms2[0].get("_stub") and bc2["stub"] == 1,
           f"stub {bc2['stub']}/2")

    return b.tong()


def _ctx(du_an):
    from eide import Config
    from eide.history import History
    from eide.ids import IdGen
    from eide.protocol.ledger import Ledger
    from eide.store import EideMd, Store
    from eide.tools import build_registry

    class C:
        config = Config.for_project(du_an)
        store = Store(config.paths.store_db)
        eide_md = EideMd.load(config.paths.eide_md, create_name="thu-cuoi")
        registry = build_registry()
        ids = IdGen(config.paths.state_dir)
        ledger = Ledger(config.paths.ledger)
        history = History(paths=config.paths, store=store, ledger=ledger, ids=ids)
        run_id = "run-thu"
        tai_lieu: dict = {}
        pending_cards: list = []
        loi_nguoi_trong_phien: list = []
        awaiting_human = False
        agent = None
        def emit(self, c): pass
    return C()


def _lam_tep(d: pathlib.Path) -> None:
    import docx

    doc = docx.Document()
    doc.add_heading("Đặc tả nội bộ", level=1)
    doc.add_heading("3.2 Electrical", level=2)
    t = doc.add_table(rows=3, cols=4)
    for j, v in enumerate(["Parameter", "Conditions", "Value", "Unit"]):
        t.rows[0].cells[j].text = v
    for j, v in enumerate(["VDD max", "TA = 25°C", "2.7 / 3.3 / 5.5", "V"]):
        t.rows[1].cells[j].text = v
    for j, v in enumerate(["Supply current", "@ 100 kHz", "12", "mA"]):
        t.rows[2].cells[j].text = v
    doc.save(str(d / "spec.docx"))

    (d / "mach.net").write_text(_NET, "utf-8")
    (d / "bom.csv").write_text("Ref,Value\nR5,10k\nU1,ATmega328P\n", "utf-8")
    (d / "bom-thieu-cot.csv").write_text("Cột A,Cột B\nx,y\n", "utf-8")
    (d / "khong-phai.net").write_text("(kicad_sch (version 1))\n", "utf-8")
    (d / "link.ld").write_text(
        "MEMORY { FLASH (rx) : ORIGIN = 0x08000000, LENGTH = 64K }\n", "utf-8")


def main() -> int:
    from eide.config import load_dotenv
    load_dotenv()
    d = (REPO / "du-lieu/thu-nghiem-cuoi").resolve()
    if d.exists():
        shutil.rmtree(d)
    d.mkdir(parents=True)
    _lam_tep(d)

    subprocess.run(["pkill", "-f", "EIDE.app/Contents/MacOS/EIDE"], check=False)
    time.sleep(1)
    subprocess.run(["defaults", "write", "vn.mobiluck.eide", "duAnPath",
                    "-string", str(d)], check=True)
    return chay(d)


if __name__ == "__main__":
    raise SystemExit(main())
