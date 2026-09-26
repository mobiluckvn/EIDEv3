#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Kiểm ING-B qua giao diện thật — bộ đọc Office.

    python tools/thu_ing_b.py

Happy: nạp `.docx` spec nội bộ → Fact ra tầng NGƯỜI có trích dẫn theo đường tiêu đề ·
nạp `.xlsx` bảng đo → trích dẫn `Sheet!ô`, công thức giữ làm nguồn của số · nạp **chính
tệp yêu cầu nâng cấp của chủ sản phẩm** (EIDE-ING-43.docx) — phép thử thật nhất có thể.

Unhappy (sáu đường): nạp Office mà không nêu nguồn · `.doc` đời cũ khi thiếu LibreOffice
· xin PDF phái sinh cho Excel · trích dẫn Word theo số trang · tài liệu nội bộ bị coi
như datasheet · nạp một định dạng không có bộ đọc.
"""

from __future__ import annotations

import pathlib
import shutil
import subprocess
import sys
import time

REPO = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "src"))
sys.path.insert(0, str(REPO / "tools"))

from thu_giao_dien import Bo, GiaoDien, PathsThu        # noqa: E402

XANH, HET = "\033[92m", "\033[0m"
EX = {"summary": "nạp tài liệu", "why": "để trích số có nguồn", "sources": [],
      "diff_prev": "—", "next": "—", "confidence": "NGUOI"}


def chay(du_an: pathlib.Path) -> int:
    b = Bo()
    g = GiaoDien(du_an)
    print("Mở app…")
    subprocess.run(["open", str(REPO / "ui/EIDEApp/EIDE.app")], check=True)
    g.san_sang(60)
    print(f"{XANH}Kênh kiểm thử giao diện đã mở{HET}\n")

    from eide.tools import build_registry
    reg = build_registry()
    ctx = _ctx(du_an)

    # ================================================================== HAPPY
    b.phan("A · .docx SPEC NỘI BỘ — TRÍCH DẪN THEO ĐƯỜNG TIÊU ĐỀ (ING01, ING06)")
    r = reg.run("doc.load", {"path": "spec-noi-bo.docx", "doc_id": "SPEC-01",
                             "nguon": "noi_bo", "explain": EX}, ctx)
    b.kiem("Nạp được .docx", r.ok, str(r.error)[:110] if not r.ok else
           f"{r.data['so_don_vi']} mục · {r.data['loai']}")
    if r.ok:
        b.kiem("Đơn vị trích dẫn là MỤC, không phải trang",
               r.data["don_vi_trich_dan"] == "mục", r.data["don_vi_trich_dan"])
        b.kiem("Trích dẫn mẫu mang đường tiêu đề",
               any(">" in m for m in r.data["trich_dan_mau"]),
               "; ".join(r.data["trich_dan_mau"])[:120])
        b.kiem("Nói rõ Word không có số trang cố định",
               "số trang" in r.data.get("note_vi", ""),
               r.data.get("note_vi", "")[:130])
        b.kiem("Tài liệu nội bộ → tầng NGƯỜI (quyết định 25/09)",
               r.data["tang_mac_dinh"] == "NGUOI", r.data["tang_mac_dinh"])

    r = reg.run("fact.extract", {"doc_id": "SPEC-01", "thuc_the": "chip:ATmega328P"}, ctx)
    b.kiem("Trích được Fact từ bảng trong Word", r.ok and r.data["so_ung_vien"] > 0,
           f"{r.data.get('so_ung_vien')} ứng viên" if r.ok else str(r.error)[:110])
    if r.ok and r.data["so_ung_vien"]:
        f = ctx.store.query_facts(limit=50)
        b.kiem("Fact ở tầng NGƯỜI, không phải BẠC",
               all(x["tier"] == "NGUOI" for x in f),
               "; ".join(sorted({x["tier"] for x in f})))
        import json as _j
        ng = [_j.loads(x["source"]) if isinstance(x["source"], str) else x["source"]
              for x in f]
        b.kiem("Mỗi Fact trỏ về đúng chỗ người mở tệp ra tìm thấy",
               all(">" in (s.get("cite") or "") for s in ng),
               (ng[0].get("cite") if ng else "")[:110])
        b.kiem("Nói rõ là nguồn nội bộ, chưa có tài liệu chuẩn",
               "nội bộ" in r.data["note_vi"], r.data["note_vi"][:120])

    b.phan("B · .xlsx BẢNG ĐO — TRÍCH DẪN Sheet!ô, GIỮ CÔNG THỨC (ING02)")
    r = reg.run("doc.load", {"path": "bang-do.xlsx", "doc_id": "DO-01",
                             "nguon": "noi_bo", "explain": EX}, ctx)
    b.kiem("Nạp được .xlsx", r.ok, f"{r.data['so_don_vi']} hàng" if r.ok else
           str(r.error)[:110])
    if r.ok:
        b.kiem("Trích dẫn là Sheet!ô",
               all("!" in m for m in r.data["trich_dan_mau"]),
               "; ".join(r.data["trich_dan_mau"]))
        tl = ctx.tai_lieu["DO-01"]
        chu = " ".join(t.chu for t in tl.trang)
        b.kiem("Công thức được giữ làm nguồn của số",
               "=AVERAGE" in chu, [x for x in chu.split("|") if "=" in x][:1])
        b.kiem("Đọc được cả hai sheet",
               len({t.trich_dan.split("!")[0] for t in tl.trang}) >= 2,
               ", ".join(sorted({t.trich_dan.split("!")[0] for t in tl.trang})))

    b.phan("C · NẠP CHÍNH TỆP YÊU CẦU CỦA CHỦ SẢN PHẨM")
    r = reg.run("ingest.file", {"path": "EIDE-ING-43.docx"}, ctx)
    b.kiem("Phân loại đúng là Word, KHÔNG phải tệp nén",
           r.ok and r.data["loai"] == "docx",
           f"{r.data.get('loai')} · {r.data.get('ly_do_phan_loai', '')[:60]}" if r.ok else "")
    r = reg.run("doc.load", {"path": "EIDE-ING-43.docx", "doc_id": "ING-43",
                             "nguon": "noi_bo", "explain": EX}, ctx)
    b.kiem("Nạp được tệp thật của anh", r.ok,
           f"{r.data['so_don_vi']} mục đọc được" if r.ok else str(r.error)[:130])
    if r.ok:
        tl = ctx.tai_lieu["ING-43"]
        chu = " ".join(t.chu for t in tl.trang)
        b.kiem("Đọc ra nội dung thật trong tài liệu",
               "ING-43" in chu or "ingest" in chu.lower() or "Office" in chu,
               chu[:130])

    b.phan("D · TÁC TỬ THẬT NẠP .docx QUA GIAO DIỆN")
    g.go("Trong dự án có spec-noi-bo.docx — đây là tài liệu nhóm mình tự viết, không "
         "phải datasheet hãng. Nạp nó vào rồi cho mình biết nó ghi VDD tối đa bao nhiêu, "
         "và số đó tin được tới đâu.")
    a = g.doi_xong(400)
    loi = a["loi_tac_tu_cuoi"]
    b.kiem("Trả lời đúng con số trong tệp", "5,5" in loi or "5.5" in loi, loi[:150])
    b.kiem("Nói rõ đây là nguồn nội bộ, chưa có tài liệu chuẩn",
           any(x in loi.lower() for x in ("nội bộ", "người", "chưa có tài liệu",
                                          "tự viết")),
           loi[:180])

    # ================================================================== UNHAPPY
    b.phan("E · SÁU ĐƯỜNG HỎNG")

    b.buoc("1. Nạp Office mà không nêu nguồn")
    spec = reg.get("doc.load")
    b.kiem("`nguon` là tham số BẮT BUỘC — tầng không được đoán",
           "nguon" in spec.params["required"],
           ", ".join(spec.params["required"]))

    b.buoc("2. .doc đời cũ")
    from eide.knowledge.office import co_libreoffice
    r = reg.run("doc.load", {"path": "datasheet-cu.doc", "doc_id": "CU-01",
                             "nguon": "nha_san_xuat", "explain": EX}, ctx)
    if co_libreoffice() is None:
        b.kiem("Thiếu LibreOffice: nói đúng lý do và đường đi tiếp, không nổ",
               not r.ok and r.error.code == "E1015"
               and "LibreOffice" in r.error.message_vi,
               r.error.message_vi[:130] if r.error else "lọt!")
        b.kiem("Gợi ý lưu lại ở định dạng mới",
               bool(r.error) and any(".docx" in x for x in r.error.alternatives),
               str(r.error.alternatives) if r.error else "")
    else:
        b.kiem("Có LibreOffice: hoặc chuyển đổi được, hoặc nói rõ vì sao không",
               r.ok or bool(r.error.message_vi),
               "chuyển được" if r.ok else r.error.message_vi[:110])

    b.buoc("3. Xin PDF phái sinh cho Excel")
    r = reg.run("doc.to_pdf", {"doc_id": "DO-01"}, ctx)
    b.kiem("Từ chối và giải thích Excel đã có trích dẫn theo ô",
           not r.ok and "không cần" in r.error.message_vi,
           r.error.message_vi[:120] if r.error else "lọt!")

    b.buoc("4. Trích dẫn Word KHÔNG được là số trang")
    tl = ctx.tai_lieu.get("SPEC-01")
    if tl:
        b.kiem("Không nhãn nào nói 'trang'",
               not any("trang" in t.trich_dan.lower() for t in tl.trang),
               "; ".join(t.trich_dan for t in tl.trang[:2]))

    b.buoc("5. Tài liệu nội bộ không được đứng ngang hàng datasheet")
    from eide.knowledge.docs import tang_mac_dinh
    b.kiem("noi_bo → NGƯỜI, nha_san_xuat → BẠC",
           tang_mac_dinh("noi_bo") == "NGUOI" and tang_mac_dinh("nha_san_xuat") == "BAC",
           f"noi_bo={tang_mac_dinh('noi_bo')} · hãng={tang_mac_dinh('nha_san_xuat')}")
    b.kiem("Nguồn lạ thì chọn phía thận trọng",
           tang_mac_dinh("khong-ro-la-gi") == "NGUOI", tang_mac_dinh("khong-ro-la-gi"))

    b.buoc("6. Nạp một định dạng chưa có bộ đọc")
    r = reg.run("doc.load", {"path": "mach.PcbDoc", "doc_id": "X",
                             "nguon": "nha_san_xuat", "explain": EX}, ctx)
    b.kiem("Từ chối đúng lý do, kèm cách xuất thay thế",
           not r.ok and r.error.code == "E1001"
           and any("netlist" in x for x in r.error.alternatives),
           f"{r.error.code} → {r.error.alternatives}" if r.error else "lọt!")

    return b.tong()


def _ctx(du_an):
    from eide import Config
    from eide.history import History
    from eide.ids import IdGen
    from eide.protocol.ledger import Ledger
    from eide.store import EideMd, Store
    from eide.tools import build_registry

    class C:
        config = Config.for_project(du_an)
        paths = PathsThu(du_an)          # sổ cái RIÊNG, kho dùng chung — xem thu_giao_dien
        store = Store(paths.store_db)
        eide_md = EideMd.load(config.paths.eide_md, create_name="thu-ing-b")
        registry = build_registry()
        ids = IdGen(paths.state_dir)
        ledger = Ledger(paths.ledger)
        history = History(paths=paths, store=store, ledger=ledger, ids=ids)
        run_id = "run-thu"
        tai_lieu: dict = {}
        pending_cards: list = []
        loi_nguoi_trong_phien: list = []
        awaiting_human = False
        def emit(self, c): pass
    return C()


def _lam_tep_mau(d: pathlib.Path) -> None:
    import docx
    import openpyxl

    doc = docx.Document()
    doc.add_heading("Đặc tả nội bộ — bộ ghi nhiệt độ", level=1)
    doc.add_paragraph("Do nhóm phần cứng viết. Chưa qua kiểm của nhà sản xuất.")
    doc.add_heading("3.2 Electrical", level=2)
    doc.add_paragraph("Thông số đo trên bo mẫu số 3, nhiệt độ phòng.")
    t = doc.add_table(rows=3, cols=4)
    for j, v in enumerate(["Parameter", "Min", "Max", "Unit"]):
        t.rows[0].cells[j].text = v
    for j, v in enumerate(["VDD max", "2.7", "5.5", "V"]):
        t.rows[1].cells[j].text = v
    for j, v in enumerate(["Supply current", "0.2", "12", "mA"]):
        t.rows[2].cells[j].text = v
    doc.add_heading("4. Ghi chú", level=2)
    doc.add_paragraph("Bo mẫu chạy ổn ở 25 °C.")
    doc.save(str(d / "spec-noi-bo.docx"))

    wb = openpyxl.Workbook()
    s1 = wb.active
    s1.title = "Bảng đo"
    s1.append(["Parameter", "Value", "Unit"])
    s1.append(["VDD max", 5.5, "V"])
    s1.append(["Supply current", 12, "mA"])
    s1["B5"] = "=AVERAGE(B2:B3)"
    s2 = wb.create_sheet("Ghi chú")
    s2.append(["Người đo", "Ngày"])
    s2.append(["Công", "2026-09-25"])
    wb.save(str(d / "bang-do.xlsx"))

    (d / "datasheet-cu.doc").write_bytes(
        b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 400)
    (d / "mach.PcbDoc").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 300)

    # Tệp THẬT của chủ sản phẩm — phép thử thật nhất có thể.
    that = REPO / "docs/20260925/EIDE-ING-43_Duong_ong_Nap_Tai_lieu_va_Trich_xuat.docx"
    if that.exists():
        shutil.copy(that, d / "EIDE-ING-43.docx")


def main() -> int:
    from eide.config import load_dotenv
    load_dotenv()
    d = (REPO / "du-lieu/thu-nghiem-ing-b").resolve()
    if d.exists():
        shutil.rmtree(d)
    d.mkdir(parents=True)
    _lam_tep_mau(d)

    subprocess.run(["pkill", "-f", "EIDE.app/Contents/MacOS/EIDE"], check=False)
    time.sleep(1)
    subprocess.run(["defaults", "write", "vn.mobiluck.eide", "duAnPath",
                    "-string", str(d)], check=True)
    return chay(d)


if __name__ == "__main__":
    raise SystemExit(main())
