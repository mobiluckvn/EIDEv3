# -*- coding: utf-8 -*-
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from ui_model import UI, REQS, all_items

F = 'Arial'
HDR = PatternFill('solid', fgColor='B8121F'); HF = Font(name=F, bold=True, color='FFFFFF', size=10)
BF = Font(name=F, size=10); BOLD = Font(name=F, bold=True, size=10)
YEL = PatternFill('solid', fgColor='FFF2CC'); GRN = PatternFill('solid', fgColor='E2F0D9'); RED = PatternFill('solid', fgColor='F8CBAD'); GRY = PatternFill('solid', fgColor='F2F2F2')
thin = Side(style='thin', color='D0D0D0'); BRD = Border(left=thin, right=thin, top=thin, bottom=thin)

def sheet(wb, name, headers, widths):
    ws = wb.create_sheet(name)
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=i, value=h); c.fill = HDR; c.font = HF; c.alignment = Alignment(wrap_text=True, vertical='center'); c.border = BRD
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.freeze_panes = 'A2'; ws.row_dimensions[1].height = 30
    return ws

def put(ws, r, vals, fills=None):
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=i, value=v); c.font = BF; c.border = BRD; c.alignment = Alignment(wrap_text=True, vertical='top')
        if fills and fills.get(i): c.fill = fills[i]

wb = Workbook(); wb.remove(wb.active)

# ---------------- Huong dan
ws = wb.create_sheet('Huong dan'); ws.column_dimensions['A'].width = 120
lines = ["MA TRẬN ÁNH XẠ HAI CHIỀU — YÊU CẦU (EIDE-MDD-40 v3.0) ↔ GIAO DIỆN (EIDE_v3.0_UI_prototype.html)", "",
         "1. Nguồn: cả HTML và Excel này được sinh từ CÙNG một mô hình (ui_model.py). Mã UI trong HTML (bật 'Chế độ truy vết') trùng với cột 'Mã UI' ở đây.",
         "2. Cấp UI: Cấp 1 = Vùng (A0–A14: thanh trạng thái, Console, 10 tab, thẻ cổng, thiết lập, dự án) · Cấp 2 = Khối (Ax.y) · Cấp 3 = Widget/Hành động (Ax.y.z).",
         "3. Sheet 'Yeu cau': 198 yêu cầu rút từ MDD-40 (nguyên tắc, bất biến, chặng, UC, HumanAct, UICommand, cổng, hiện vật, lớp giải thích, quy tắc trình bày, đồng bộ khi sửa, changeset, snapshot, tầng, luồng datasheet, bố cục, ngữ cảnh, bộ nhớ, sự cố, năng lực). Cột 'Số UI ánh xạ' là CÔNG THỨC đếm từ sheet 'Anh xa'; = 0 → THIẾU.",
         "4. Sheet 'UI 3 cap': toàn bộ cây UI. Cột 'Số yêu cầu' là công thức; widget cấp 3 = 0 → THIẾU (khối cấp 2 và vùng cấp 1 được phủ qua widget con hoặc yêu cầu gán trực tiếp).",
         "5. Sheet 'Anh xa' (dạng dài): mỗi dòng một cặp (UI, yêu cầu). Đây là bảng gốc cho hai chiều; sửa ở đây thì hai sheet kiểm tra tự cập nhật.",
         "6. Sheet 'REQ -> UI' và 'UI -> REQ': bảng đọc nhanh, mỗi dòng một ID với danh sách ID phía kia (chuỗi do trình sinh ghép) + công thức đếm để đối chiếu.",
         "7. Sheet 'Kiem tra': tổng hợp; ô 'KẾT LUẬN' phải là ĐẠT (0 yêu cầu chưa phủ, 0 widget chưa ánh xạ).",
         "8. Chỉ điền cột nền vàng (Ghi chú rà soát, Trạng thái hiện thực) khi dùng để theo dõi hiện thực; các cột khác do trình sinh quản lý.",
         "", "Ngày sinh: 25/09/2026 · Nguồn yêu cầu: EIDE-MDD-40 v3.0 · Tác giả: Vũ Trí Công (Đề án ThS Kỹ thuật Điện tử — PTIT, GVHD TS. Nguyễn Trung Hiếu)"]
for i, l in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=l); c.font = Font(name=F, bold=(i == 1), size=11 if i == 1 else 10); c.alignment = Alignment(wrap_text=True)

# ---------------- Anh xa (long) — build first for formulas
pairs = []  # (ui_id, ui_name, level, req)
for a in UI:
    for b in a['blocks']:
        for r in b['reqs']: pairs.append((b['id'], b['name'], 2, r))
        for w in b['items']:
            for r in w['reqs']: pairs.append((w['id'], w['name'], 3, r))
REQ = {r['id']: r for r in REQS}
wsm = sheet(wb, 'Anh xa', ['Mã UI', 'Tên UI', 'Cấp', 'Mã yêu cầu', 'Nhóm yêu cầu', 'Mô tả yêu cầu', 'Nguồn §', 'Ghi chú rà soát'], [12, 60, 6, 16, 22, 60, 18, 30])
for i, (u, n, lv, r) in enumerate(pairs, 2):
    put(wsm, i, [u, n, lv, r, REQ[r]['group'], REQ[r]['desc'], REQ[r]['src'], ''], {8: YEL})
NP = len(pairs) + 1
RNG_UI = f"'Anh xa'!$A$2:$A${NP}"; RNG_RQ = f"'Anh xa'!$D$2:$D${NP}"

# ---------------- Yeu cau
wsr = sheet(wb, 'Yeu cau', ['Mã yêu cầu', 'Nhóm', 'Mô tả', 'Nguồn §', 'Số UI ánh xạ (công thức)', 'Trạng thái phủ', 'Trạng thái hiện thực', 'Ghi chú rà soát'], [14, 24, 70, 18, 14, 14, 18, 30])
dv = DataValidation(type='list', formula1='"Chưa làm,Đang làm,Đã làm,Không áp dụng"', allow_blank=True); wsr.add_data_validation(dv)
for i, r in enumerate(REQS, 2):
    put(wsr, i, [r['id'], r['group'], r['desc'], r['src'], f'=COUNTIF({RNG_RQ},A{i})', f'=IF(E{i}>0,"Đã phủ","THIẾU")', 'Chưa làm', ''], {7: YEL, 8: YEL})
    dv.add(f'G{i}')
NR = len(REQS) + 1
from openpyxl.formatting.rule import CellIsRule, FormulaRule
wsr.conditional_formatting.add(f'F2:F{NR}', CellIsRule(operator='equal', formula=['"THIẾU"'], fill=RED))
wsr.conditional_formatting.add(f'F2:F{NR}', CellIsRule(operator='equal', formula=['"Đã phủ"'], fill=GRN))

# ---------------- UI 3 cap
wsu = sheet(wb, 'UI 3 cap', ['Mã UI', 'Cấp', 'Vùng (cấp 1)', 'Khối (cấp 2)', 'Widget / hành động (cấp 3)', 'Loại', 'Hiện vật', 'Số yêu cầu (công thức)', 'Trạng thái ánh xạ', 'Trạng thái hiện thực', 'Ghi chú rà soát'], [12, 6, 26, 34, 60, 10, 22, 12, 14, 18, 30])
dv2 = DataValidation(type='list', formula1='"Chưa làm,Đang làm,Đã làm"', allow_blank=True); wsu.add_data_validation(dv2)
row = 2
for a in UI:
    # cấp 1: đếm mọi cặp có UI bắt đầu bằng "A?."
    put(wsu, row, [a['id'], 1, a['name'], '', '', a['kind'], '', f'=COUNTIF({RNG_UI},"{a["id"]}.*")', f'=IF(H{row}>0,"Đã ánh xạ","THIẾU")', 'Chưa làm', ''], {1: GRY, 2: GRY, 3: GRY, 10: YEL, 11: YEL}); dv2.add(f'J{row}'); row += 1
    for b in a['blocks']:
        put(wsu, row, [b['id'], 2, a['name'], b['name'], '', 'khối', b.get('art', ''), f'=COUNTIF({RNG_UI},"{b["id"]}")+COUNTIF({RNG_UI},"{b["id"]}.*")', f'=IF(H{row}>0,"Đã ánh xạ","THIẾU")', 'Chưa làm', ''], {10: YEL, 11: YEL}); dv2.add(f'J{row}'); row += 1
        for w in b['items']:
            put(wsu, row, [w['id'], 3, a['name'], b['name'], w['name'], w['kind'], b.get('art', ''), f'=COUNTIF({RNG_UI},A{row})', f'=IF(H{row}>0,"Đã ánh xạ","THIẾU")', 'Chưa làm', ''], {10: YEL, 11: YEL}); dv2.add(f'J{row}'); row += 1
NU = row - 1
wsu.conditional_formatting.add(f'I2:I{NU}', CellIsRule(operator='equal', formula=['"THIẾU"'], fill=RED))
wsu.conditional_formatting.add(f'I2:I{NU}', CellIsRule(operator='equal', formula=['"Đã ánh xạ"'], fill=GRN))

# ---------------- REQ -> UI
used = {}
for u, n, lv, r in pairs: used.setdefault(r, []).append(u)
ws1 = sheet(wb, 'REQ -> UI', ['Mã yêu cầu', 'Nhóm', 'Mô tả', 'Danh sách mã UI (ghép)', 'Số UI (ghép)', 'Số UI (công thức từ Anh xa)', 'Khớp?'], [14, 22, 60, 60, 10, 12, 8])
for i, r in enumerate(REQS, 2):
    ids = sorted(set(used.get(r['id'], [])))
    put(ws1, i, [r['id'], r['group'], r['desc'], ', '.join(ids), len(ids), f'=SUMPRODUCT(({RNG_RQ}=A{i})*1)', f'=IF(AND(F{i}>0,F{i}=E{i}),"✓","THIẾU")'])
# ---------------- UI -> REQ
byui = {}
for u, n, lv, r in pairs: byui.setdefault(u, (n, lv, []))[2].append(r)
ws2 = sheet(wb, 'UI -> REQ', ['Mã UI', 'Cấp', 'Tên UI', 'Danh sách mã yêu cầu (ghép)', 'Số yêu cầu (ghép)', 'Số yêu cầu (công thức từ Anh xa)', 'Khớp?'], [12, 6, 60, 60, 10, 12, 8])
r_i = 2
for a in UI:
    for b in a['blocks']:
        for u in [b['id']] + [w['id'] for w in b['items']]:
            if u not in byui: continue
            n, lv, rs = byui[u]; rs = sorted(set(rs))
            put(ws2, r_i, [u, lv, n, ', '.join(rs), len(rs), f'=SUMPRODUCT(({RNG_UI}=A{r_i})*1)', f'=IF(F{r_i}>0,"✓","THIẾU")']); r_i += 1
NU2 = r_i - 1

# ---------------- Kiem tra
wk = sheet(wb, 'Kiem tra', ['Chỉ số', 'Giá trị', 'Ghi chú'], [50, 14, 60])
rows = [
    ('Tổng số yêu cầu (MDD-40)', f"=COUNTA('Yeu cau'!A2:A{NR})", 'Sheet Yeu cau'),
    ('Yêu cầu CHƯA được UI nào phủ', f"=COUNTIF('Yeu cau'!F2:F{NR},\"THIẾU\")", 'Phải = 0 (chiều Yêu cầu → UI)'),
    ('Tổng số vùng (cấp 1)', f"=COUNTIF('UI 3 cap'!B2:B{NU},1)", ''),
    ('Tổng số khối (cấp 2)', f"=COUNTIF('UI 3 cap'!B2:B{NU},2)", ''),
    ('Tổng số widget/hành động (cấp 3)', f"=COUNTIF('UI 3 cap'!B2:B{NU},3)", ''),
    ('Widget cấp 3 CHƯA ánh xạ yêu cầu', f"=COUNTIFS('UI 3 cap'!B2:B{NU},3,'UI 3 cap'!I2:I{NU},\"THIẾU\")", 'Phải = 0 (chiều UI → Yêu cầu)'),
    ('Khối cấp 2 CHƯA ánh xạ', f"=COUNTIFS('UI 3 cap'!B2:B{NU},2,'UI 3 cap'!I2:I{NU},\"THIẾU\")", 'Phải = 0'),
    ('Vùng cấp 1 CHƯA ánh xạ', f"=COUNTIFS('UI 3 cap'!B2:B{NU},1,'UI 3 cap'!I2:I{NU},\"THIẾU\")", 'Phải = 0'),
    ('Tổng số cặp ánh xạ', f"=COUNTA('Anh xa'!A2:A{NP})", ''),
    ('Cặp có mã yêu cầu không tồn tại', f"=SUMPRODUCT((COUNTIF('Yeu cau'!$A$2:$A${NR},'Anh xa'!D2:D{NP})=0)*1)", 'Phải = 0'),
    ('KẾT LUẬN', '=IF(AND(B3=0,B7=0,B8=0,B9=0,B11=0),"ĐẠT — kín hai chiều","KHÔNG ĐẠT — xem ô đỏ")', ''),
    ('Yêu cầu theo trạng thái hiện thực: Đã làm', f"=COUNTIF('Yeu cau'!G2:G{NR},\"Đã làm\")", 'Cột vàng do người điền'),
    ('Yêu cầu theo trạng thái hiện thực: Đang làm', f"=COUNTIF('Yeu cau'!G2:G{NR},\"Đang làm\")", ''),
    ('Yêu cầu theo trạng thái hiện thực: Chưa làm', f"=COUNTIF('Yeu cau'!G2:G{NR},\"Chưa làm\")", ''),
]
for i, (a_, b_, c_) in enumerate(rows, 2):
    put(wk, i, [a_, b_, c_]); wk.cell(row=i, column=1).font = BOLD if 'KẾT LUẬN' in a_ else BF
wk.conditional_formatting.add('B3', CellIsRule(operator='greaterThan', formula=['0'], fill=RED))
for c in ['B7', 'B8', 'B9', 'B11']: wk.conditional_formatting.add(c, CellIsRule(operator='greaterThan', formula=['0'], fill=RED))
wk.conditional_formatting.add('B12', FormulaRule(formula=['LEFT(B12,3)="ĐẠT"'], fill=GRN))

# ---------------- Nhom yeu cau (thống kê theo nhóm)
groups = sorted(set(r['group'] for r in REQS), key=lambda g: [r['group'] for r in REQS].index(g))
wg = sheet(wb, 'Theo nhom', ['Nhóm yêu cầu', 'Số yêu cầu', 'Số yêu cầu đã phủ', 'Số cặp ánh xạ'], [30, 12, 16, 14])
for i, g in enumerate(groups, 2):
    put(wg, i, [g, f"=COUNTIF('Yeu cau'!B2:B{NR},A{i})", f"=COUNTIFS('Yeu cau'!B2:B{NR},A{i},'Yeu cau'!F2:F{NR},\"Đã phủ\")", f"=COUNTIF('Anh xa'!E2:E{NP},A{i})"])

wb.move_sheet('Kiem tra', offset=-(len(wb.sheetnames) - 2))
wb.save(str(__import__('pathlib').Path(__file__).resolve().parent / 'EIDE_v3.0_Anh_xa_Yeu_cau_UI.xlsx'))
print('xlsx ok', len(pairs), 'cặp')
