# -*- coding: utf-8 -*-
"""Bộ đọc Office — EIDE-ING-43 §4.2, §6. Bước ING-B.

Ca đo: ING01 (.docx có bảng), ING02 (.xlsx hai sheet), ING03 (.doc cũ), ING-06 (trích
dẫn theo loại), ING-19 (tầng theo nguồn).

Điều cả bộ này canh không phải "đọc được chữ không" — thư viện làm việc đó. Nó canh câu
**"số này ở đâu trong tệp?"**: một Fact không chỉ được về đúng chỗ người mở tệp ra tìm
thấy thì N1 chỉ còn là một lời hứa.
"""

from __future__ import annotations

import shutil

import pytest

from eide.knowledge.docs import TANG_THEO_NGUON, tang_mac_dinh
from eide.knowledge.office import (co_libreoffice, doc_docx, doc_pptx, doc_xlsx,
                                   la_bang_thong_so, nap_office)


# =========================================================================== dựng tệp mẫu
@pytest.fixture
def tep_docx(tmp_path):
    import docx

    d = docx.Document()
    d.add_heading("Đặc tả nội bộ bộ ghi nhiệt", level=1)
    d.add_paragraph("Tài liệu này do nhóm phần cứng viết, chưa qua kiểm của nhà sản xuất.")
    d.add_heading("3.2 Electrical", level=2)
    d.add_paragraph("Các thông số điện đo trên bo mẫu số 3.")

    t = d.add_table(rows=3, cols=4)
    for j, v in enumerate(["Parameter", "Min", "Max", "Unit"]):
        t.rows[0].cells[j].text = v
    for j, v in enumerate(["VDD max", "2.7", "5.5", "V"]):
        t.rows[1].cells[j].text = v
    for j, v in enumerate(["Supply current", "0.2", "12", "mA"]):
        t.rows[2].cells[j].text = v

    d.add_heading("4. Ghi chú", level=2)
    d.add_paragraph("Bo mẫu số 3 chạy ổn ở 25 °C.")
    p = tmp_path / "spec-noi-bo.docx"
    d.save(str(p))
    return p


@pytest.fixture
def tep_xlsx(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    s1 = wb.active
    s1.title = "Bảng đo"
    s1.append(["Parameter", "Value", "Unit"])
    s1.append(["VDD max", 5.5, "V"])
    s1.append(["Supply current", 12, "mA"])
    s1["B4"] = "=AVERAGE(B2:B3)"

    s2 = wb.create_sheet("Ghi chú")
    s2.append(["Người đo", "Ngày"])
    s2.append(["Công", "2026-09-25"])

    p = tmp_path / "bang-do.xlsx"
    wb.save(str(p))
    return p


@pytest.fixture
def tep_pptx(tmp_path):
    from pptx import Presentation

    pr = Presentation()
    sl = pr.slides.add_slide(pr.slide_layouts[1])
    sl.shapes.title.text = "Kiến trúc bộ ghi nhiệt"
    sl.placeholders[1].text = "ATmega328P · thẻ SD · MTP"
    p = tmp_path / "trinh-bay.pptx"
    pr.save(str(p))
    return p


# =========================================================================== ING01 docx
def test_ING01_docx_doc_duoc_doan_va_bang(tep_docx):
    tl = doc_docx(tep_docx, doc_id="SPEC-01")
    assert tl.loai == "docx" and tl.don_vi_trich_dan == "mục"
    chu = " ".join(t.chu for t in tl.trang)
    assert "VDD max" in chu and "5.5" in chu
    assert "bo mẫu số 3" in chu.lower()


def test_ING06_trich_dan_docx_la_DUONG_TIEU_DE_khong_phai_so_trang(tep_docx):
    """Word không có số trang cố định — trích dẫn theo trang là một lời hứa sai."""
    tl = doc_docx(tep_docx, doc_id="SPEC-01")
    dong_bang = [t for t in tl.trang if "VDD max" in t.chu]
    assert dong_bang, "phải đọc được hàng bảng"
    nhan = dong_bang[0].trich_dan
    assert "3.2 Electrical" in nhan, nhan
    assert "Bảng 1" in nhan and "dòng 2" in nhan, nhan
    assert "trang" not in nhan.lower()


def test_ING01_bang_thong_so_phan_biet_voi_bang_trinh_bay(tep_docx):
    assert la_bang_thong_so(["Parameter", "Min", "Max", "Unit"])
    assert la_bang_thong_so(["Thông số", "Giá trị", "Đơn vị"])
    assert not la_bang_thong_so(["Tên", "Ảnh"])
    assert not la_bang_thong_so([])

    tl = doc_docx(tep_docx, doc_id="SPEC-01")
    assert not any("bảng trình bày" in t.trich_dan for t in tl.trang), \
        "bảng có Parameter/Min/Max phải được nhận là bảng thông số"


def test_docx_giu_cay_tieu_de_theo_MUC(tep_docx):
    tl = doc_docx(tep_docx, doc_id="SPEC-01")
    ghi_chu = [t for t in tl.trang if "25 °C" in t.chu]
    assert ghi_chu and "4. Ghi chú" in ghi_chu[0].trich_dan
    # H2 mới phải THAY H2 cũ, không nối thêm.
    assert "3.2 Electrical" not in ghi_chu[0].trich_dan


# =========================================================================== ING02 xlsx
def test_ING02_xlsx_moi_sheet_deu_doc(tep_xlsx):
    tl = doc_xlsx(tep_xlsx, doc_id="DO-01")
    assert tl.loai == "xlsx" and tl.don_vi_trich_dan == "ô"
    nhan = [t.trich_dan for t in tl.trang]
    assert any(n.startswith("Bảng đo!") for n in nhan)
    assert any(n.startswith("Ghi chú!") for n in nhan)


def test_ING02_trich_dan_la_Sheet_o(tep_xlsx):
    tl = doc_xlsx(tep_xlsx, doc_id="DO-01")
    vdd = [t for t in tl.trang if "VDD max" in t.chu]
    assert vdd and vdd[0].trich_dan == "Bảng đo!A2", vdd[0].trich_dan


def test_ING02_cong_thuc_duoc_GIU_lam_nguon_cua_so(tep_xlsx):
    """`=AVERAGE(B2:B3)` nói với người rà soát nhiều hơn con số nó tính ra."""
    tl = doc_xlsx(tep_xlsx, doc_id="DO-01")
    chu = " ".join(t.chu for t in tl.trang)
    assert "=AVERAGE(B2:B3)" in chu


# =========================================================================== pptx
def test_pptx_trich_dan_theo_slide(tep_pptx):
    tl = doc_pptx(tep_pptx, doc_id="TB-01")
    assert tl.loai == "pptx" and tl.don_vi_trich_dan == "slide"
    assert tl.trang and tl.trang[0].trich_dan == "slide 1"
    assert "ATmega328P" in tl.trang[0].chu


# =========================================================================== ING-19 tầng
def test_ING19_office_noi_bo_la_tang_NGUOI():
    """Quyết định chủ sản phẩm 25/09: tài liệu tự viết → NGƯỜI, không phải BẠC."""
    assert tang_mac_dinh("noi_bo") == "NGUOI"
    assert tang_mac_dinh("nha_san_xuat") == "BAC"
    assert tang_mac_dinh("ben_thu_ba") == "BAC"


def test_nguon_la_thi_chon_phia_THAN_TRONG():
    assert tang_mac_dinh("khong_biet_la_gi") == "NGUOI"
    assert set(TANG_THEO_NGUON) == {"nha_san_xuat", "ben_thu_ba", "noi_bo"}


def test_fact_tu_tai_lieu_noi_bo_noi_ro_chua_co_tai_lieu_chuan(tep_docx):
    from eide.knowledge.docs import FactUngVien, fact_tu_ung_vien

    tl = doc_docx(tep_docx, doc_id="SPEC-01")
    uv = FactUngVien(khoa="vdd.max", gia_tri=5.5, don_vi="V", trang=tl.trang[0].so,
                     trich_doan="VDD max | 2.7 | 5.5 | V", thuc_the="chip:X",
                     nguyen_van="5.5 V")
    f = fact_tu_ung_vien(uv, doc=tl, tier="NGUOI")
    assert f["tier"] == "NGUOI"
    assert "tài liệu chuẩn" in f["explain"]["next"]
    # Trích dẫn phải là nhãn thật, không phải "tr.N".
    assert "tr." not in f["explain"]["sources"][0]["ref"]
    assert f["source"]["cite"] == tl.trang[0].trich_dan


# =========================================================================== ING03 đời cũ
def test_ING03_doc_cu_can_LibreOffice_va_noi_that_khi_thieu(tmp_path):
    p = tmp_path / "cu.doc"
    p.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 300)
    tl, vi_sao = nap_office(p, loai="office_cu", doc_id="X",
                            thu_muc_tam=tmp_path / "tam")
    if co_libreoffice() is None:
        assert tl is None
        assert "LibreOffice" in vi_sao and "định dạng mới" in vi_sao, vi_sao
    else:
        # Có LibreOffice: tệp giả này không phải .doc thật nên vẫn hỏng — nhưng phải
        # hỏng có lý do đọc được, không phải một traceback.
        assert tl is None or tl.chuyen_doi_tu == ".doc"
        assert tl is not None or vi_sao


def test_nap_office_khong_nem_ngoai_le_khi_khong_doc_duoc(tmp_path):
    """Thiếu LibreOffice là tình huống bình thường trên máy người dùng, không phải sự cố."""
    p = tmp_path / "la.xyz"
    p.write_bytes(b"rac")
    tl, vi_sao = nap_office(p, loai="office_cu", doc_id="X",
                            thu_muc_tam=tmp_path / "tam")
    assert tl is None and vi_sao


def test_loai_la_thi_noi_ro(tmp_path):
    p = tmp_path / "x.bin"
    p.write_bytes(b"x")
    tl, vi_sao = nap_office(p, loai="khong_biet", doc_id="X",
                            thu_muc_tam=tmp_path / "tam")
    assert tl is None and "khong_biet" in vi_sao


# =========================================================================== qua công cụ
def _ctx(agent):
    from eide.loop import TurnContext

    return TurnContext(config=agent.config, store=agent.store, ledger=agent.ledger,
                       eide_md=agent.eide_md, ids=agent.ids, registry=agent.registry,
                       emit=lambda c: None, history=agent.history, run_id="run-1")


_EX = {"summary": "nạp", "why": "để trích số", "sources": [], "diff_prev": "—",
       "next": "—", "confidence": "NGUOI"}


def test_doc_load_nhan_docx_va_bat_buoc_neu_nguon(make_agent, tep_docx):
    agent = make_agent([])
    ctx = _ctx(agent)
    shutil.copy(tep_docx, agent.config.paths.project_root / "spec.docx")

    spec = agent.registry.get("doc.load")
    assert "nguon" in spec.params["required"], "tầng phụ thuộc nguồn — không được đoán"

    r = agent.registry.run("doc.load", {"path": "spec.docx", "doc_id": "SPEC-01",
                                        "nguon": "noi_bo", "explain": _EX}, ctx)
    assert r.ok, getattr(r.error, "message_vi", "")
    assert r.data["loai"] == "docx" and r.data["tang_mac_dinh"] == "NGUOI"
    assert r.data["don_vi_trich_dan"] == "mục"
    assert "số trang" in r.data["note_vi"] and "nội bộ" in r.data["note_vi"]


def test_fact_extract_dung_tang_theo_NGUON_cua_tai_lieu(make_agent, tep_docx):
    agent = make_agent([])
    ctx = _ctx(agent)
    shutil.copy(tep_docx, agent.config.paths.project_root / "spec.docx")
    agent.registry.run("doc.load", {"path": "spec.docx", "doc_id": "SPEC-01",
                                    "nguon": "noi_bo", "explain": _EX}, ctx)
    r = agent.registry.run("fact.extract",
                           {"doc_id": "SPEC-01", "thuc_the": "chip:ATmega328P"}, ctx)
    assert r.ok
    assert r.data["tang"] == "NGUOI" and r.data["nguon"] == "noi_bo"
    if r.data["so_ung_vien"]:
        f = agent.store.query_facts(limit=50)
        assert all(x["tier"] == "NGUOI" for x in f)
        assert "nội bộ" in r.data["note_vi"]


def test_doc_load_xlsx_trich_dan_theo_o(make_agent, tep_xlsx):
    agent = make_agent([])
    ctx = _ctx(agent)
    shutil.copy(tep_xlsx, agent.config.paths.project_root / "do.xlsx")
    r = agent.registry.run("doc.load", {"path": "do.xlsx", "doc_id": "DO-01",
                                        "nguon": "noi_bo", "explain": _EX}, ctx)
    assert r.ok and r.data["don_vi_trich_dan"] == "ô"
    assert any("!" in m for m in r.data["trich_dan_mau"]), r.data["trich_dan_mau"]


def test_doc_to_pdf_tu_choi_tai_lieu_khong_can(make_agent, tep_xlsx):
    agent = make_agent([])
    ctx = _ctx(agent)
    shutil.copy(tep_xlsx, agent.config.paths.project_root / "do.xlsx")
    agent.registry.run("doc.load", {"path": "do.xlsx", "doc_id": "DO-01",
                                    "nguon": "noi_bo", "explain": _EX}, ctx)
    r = agent.registry.run("doc.to_pdf", {"doc_id": "DO-01"}, ctx)
    assert not r.ok and "không cần" in r.error.message_vi


def test_doc_load_van_tu_choi_dinh_dang_khong_doc_duoc(make_agent):
    agent = make_agent([])
    ctx = _ctx(agent)
    (agent.config.paths.project_root / "mach.PcbDoc").write_bytes(
        b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 200)
    r = agent.registry.run("doc.load", {"path": "mach.PcbDoc", "doc_id": "X",
                                        "nguon": "nha_san_xuat", "explain": _EX}, ctx)
    assert not r.ok and r.error.code == "E1001"
    assert any("netlist" in x for x in r.error.alternatives)


def test_o_CHI_CO_cong_thuc_khong_bi_bien_mat(tmp_path):
    """Tệp chưa từng mở bằng Excel thì ô công thức không có giá trị đã tính.

    `data_only=True` trả None, cả hàng bị coi là rỗng, và một dòng dữ liệu biến mất mà
    không ai biết. Thà hiện `=AVERAGE(...)` còn hơn không hiện gì.
    """
    import openpyxl

    wb = openpyxl.Workbook()
    sh = wb.active
    sh.title = "Đo"
    sh["A1"] = "Trung bình"
    sh["B1"] = "=AVERAGE(B2:B9)"
    p = tmp_path / "ct.xlsx"
    wb.save(str(p))

    tl = doc_xlsx(p, doc_id="X")
    chu = " ".join(t.chu for t in tl.trang)
    assert "=AVERAGE(B2:B9)" in chu and "Trung bình" in chu


def test_ING01_trich_so_tu_BANG_khi_don_vi_o_cot_rieng(tep_docx):
    """Datasheet đặt đơn vị ở cột riêng: "VDD max | 2.7 | 5.5 | V".

    Nối cả hàng thành chuỗi rồi tìm "số kèm đơn vị" thì `5.5` và `V` cách nhau một dấu
    gạch — bộ trích trả về RỖNG trên đúng loại tài liệu nó sinh ra để đọc.
    """
    from eide.knowledge.docs import trich_fact_ung_vien

    tl = doc_docx(tep_docx, doc_id="SPEC-01")
    uv = trich_fact_ung_vien(tl, thuc_the="chip:ATmega328P")
    assert uv, "không trích được gì từ bảng Word"

    theo_khoa = {u.khoa: u for u in uv}
    assert "vdd.max" in theo_khoa, sorted(theo_khoa)
    assert theo_khoa["vdd.max"].gia_tri == 5.5
    assert theo_khoa["vdd.max"].don_vi == "V"
    assert "vdd.min" in theo_khoa and theo_khoa["vdd.min"].gia_tri == 2.7
    # Trích dẫn vẫn phải là đường tiêu đề, không phải số trang.
    assert "3.2 Electrical" in tl.trich_dan(theo_khoa["vdd.max"].trang)


def test_hang_bang_khong_co_ten_thong_so_thi_bo_qua(tmp_path):
    import docx

    d = docx.Document()
    t = d.add_table(rows=2, cols=3)
    for j, v in enumerate(["Parameter", "Max", "Unit"]):
        t.rows[0].cells[j].text = v
    for j, v in enumerate(["Màu vỏ hộp", "3", "cái"]):
        t.rows[1].cells[j].text = v
    p = tmp_path / "la.docx"
    d.save(str(p))

    from eide.knowledge.docs import trich_fact_ung_vien
    tl = doc_docx(p, doc_id="X")
    assert trich_fact_ung_vien(tl, thuc_the="chip:X") == [], \
        "tên cột không khớp thông số nào thì KHÔNG được bịa ra Fact"


# ============================== bảng BẢN ĐỒ CHÂN — tài liệu bàn giao phần cứng (ING-43 §5)
@pytest.fixture
def tep_ban_giao(tmp_path):
    """Một tài liệu bàn giao phần cứng thu nhỏ, đúng hình dạng tài liệu thật của dự án."""
    import docx

    d = docx.Document()
    d.add_heading("4. Bản đồ chân", level=1)
    d.add_heading("4.2 Bảng chân đầy đủ", level=2)
    t = d.add_table(rows=4, cols=4)
    for j, v in enumerate(["Chân", "Hướng", "Net · khối", "Chức năng và ghi chú"]):
        t.rows[0].cells[j].text = v
    for i, r in enumerate((
            ["D4 · PD4", "ra", "DIR1 · A4988 #1", "Chiều quay bánh PHẢI. Mức THẤP = tiến."],
            ["A4 · PC4", "hai chiều", "SDA · MPU6050", "TWI SDA, 400 kHz."],
            ["D13 · PB5", "ra", "—", "Không có net; điểm đo thời gian."]), start=1):
        for j, v in enumerate(r):
            t.rows[i].cells[j].text = v
    p = tmp_path / "ban-giao.docx"
    d.save(str(p))
    return p


def test_bang_ban_do_chan_duoc_nhan_la_BANG(tep_ban_giao):
    """Từ điển cột cũ chỉ biết datasheet điện (parameter/min/max/unit), nên bảng
    `Chân | Hướng | Net | Chức năng` bị đọc như một dòng chữ — và cả bản đồ chân của một
    tài liệu BÀN GIAO PHẦN CỨNG biến mất khỏi phần trích xuất."""
    tl = doc_docx(tep_ban_giao, doc_id="BG")
    hang = [t for t in tl.trang if t.loai_bang == "chan"]
    assert len(hang) == 3, [t.o for t in tl.trang if t.o]
    assert hang[0].cot[0] == "Chân"


def test_bang_chan_KHONG_bi_nham_voi_bang_thong_so(tep_ban_giao, tep_docx):
    from eide.knowledge.office import la_bang_chan, la_bang_thong_so

    assert la_bang_chan(["Chân", "Hướng", "Net · khối", "Chức năng"]) is True
    assert la_bang_thong_so(["Chân", "Hướng", "Net · khối", "Chức năng"]) is False
    assert la_bang_chan(["Parameter", "Min", "Max", "Unit"]) is False
    # Bảng chỉ có mô tả thì KHÔNG phải bản đồ chân — thiếu cột "chân" thì không có gì để gán.
    assert la_bang_chan(["Chức năng", "Ghi chú"]) is False


def test_trich_chan_tach_ten_bo_va_ten_cong_va_net(tep_ban_giao):
    from eide.knowledge.docs import trich_chan_ung_vien

    tl = doc_docx(tep_ban_giao, doc_id="BG")
    uv = {x.so_chan: x for x in trich_chan_ung_vien(tl)}
    assert set(uv) == {"D4", "A4", "D13"}
    # Giữ CẢ HAI tên: mất tên cổng thì không viết được DDR, mất tên bo thì người cầm bo
    # không tìm ra chân.
    assert uv["D4"].cong == "PD4" and uv["D4"].net == "DIR1" and uv["D4"].khoi == "A4988 #1"
    assert uv["D4"].huong == "ra" and "Mức THẤP = tiến" in uv["D4"].chuc_nang
    assert uv["A4"].huong == "hai_chieu" and uv["A4"].net == "SDA"
    # Chân không có net vẫn phải vào danh sách — nó là một chân có thật trên bo.
    assert uv["D13"].net in ("", "—")


def test_moi_chan_mang_theo_TRICH_DAN_toi_dung_dong_bang(tep_ban_giao):
    from eide.knowledge.docs import fact_tu_chan, trich_chan_ung_vien

    tl = doc_docx(tep_ban_giao, doc_id="BG")
    uv = trich_chan_ung_vien(tl)
    fs = fact_tu_chan(uv[0], doc=tl, chip="ATmega328P")
    assert {f["key"] for f in fs} >= {"ten", "net", "huong"}
    assert all(f["subject"] == "pin:ATmega328P.D4" for f in fs)
    cite = fs[0]["source"]["cite"]
    assert "4.2 Bảng chân đầy đủ" in cite and "dòng" in cite, cite
    # Khoá `ten` phải là TÊN CỔNG, vì đó là thứ `ckm.chan_tu_fact` và mã thanh ghi dùng.
    assert next(f for f in fs if f["key"] == "ten")["value"] == "PD4"


def test_fact_chan_di_THANG_vao_cay_khoi_khong_phai_chep_tay(tep_ban_giao):
    """Chép tay bản đồ chân từ tài liệu sang bản đồ mạch là chỗ sai không ai kiểm được."""
    from eide.knowledge.ckm import chan_tu_fact
    from eide.knowledge.docs import fact_tu_chan, trich_chan_ung_vien

    tl = doc_docx(tep_ban_giao, doc_id="BG")
    fs = [f for x in trich_chan_ung_vien(tl)
          for f in fact_tu_chan(x, doc=tl, chip="ATmega328P")]
    chan = chan_tu_fact(fs, "ATmega328P")
    assert set(chan) == {"D4", "A4", "D13"}
    # Tầng mặc định của `fact_tu_chan` là BẠC: nguồn đã duyệt, DÒNG chưa ai xác nhận.
    assert chan["D4"].ten == "PD4" and chan["D4"].tier == "BAC"


def test_mo_lai_du_an_thi_tai_lieu_van_DOC_DUOC(make_agent, tep_ban_giao):
    """Đóng app rồi mở lại: hiện vật `doc` còn trong kho, nhưng tài liệu đã phân tích thì
    không — nó nằm trong bộ nhớ tiến trình. Không có đường đọc lại thì mọi công cụ đọc tài
    liệu trả "chưa được nạp" cho đúng tài liệu mà tab Tài liệu đang hiện."""
    agent = make_agent([])
    ctx = _ctx(agent)
    shutil.copy(tep_ban_giao, agent.config.paths.project_root / "bg.docx")
    agent.registry.run("doc.load", {"path": "bg.docx", "doc_id": "BG-1",
                                    "nguon": "noi_bo", "explain": _EX}, ctx)
    assert agent.registry.run(
        "fact.extract_pinout", {"doc_id": "BG-1", "chip": "ATmega328P"}, ctx).ok

    ctx.tai_lieu.clear()                       # đúng thứ xảy ra khi mở lại dự án
    r = agent.registry.run("fact.extract_pinout",
                           {"doc_id": "BG-1", "chip": "ATmega328P"}, ctx)
    assert r.ok, getattr(r.error, "message_vi", "")
    assert r.data["so_chan"] == 3


def test_tep_DOI_tu_lan_nap_thi_NOI_RA_chu_khong_dung_am_tham(make_agent, tep_ban_giao):
    """Trích dẫn đã ghi trỏ vào bản CŨ. Đọc bản mới dưới tên cũ là làm mọi trích dẫn sai đi
    mà không ai phát hiện được."""
    import docx

    agent = make_agent([])
    ctx = _ctx(agent)
    p = agent.config.paths.project_root / "bg.docx"
    shutil.copy(tep_ban_giao, p)
    agent.registry.run("doc.load", {"path": "bg.docx", "doc_id": "BG-2",
                                    "nguon": "noi_bo", "explain": _EX}, ctx)
    ctx.tai_lieu.clear()

    d = docx.Document(str(p))
    d.add_paragraph("Sửa tay sau khi bàn giao.")
    d.save(str(p))

    r = agent.registry.run("fact.extract_pinout",
                           {"doc_id": "BG-2", "chip": "ATmega328P"}, ctx)
    assert not r.ok and r.error.code == "E2005", getattr(r, "data", None)
    assert "đã ĐỔI kể từ lần nạp" in r.error.message_vi


def test_doc_read_doc_duoc_MUC_va_BANG_kem_trich_dan(make_agent, tep_ban_giao):
    """Trước khi có `doc.read`, tác tử nạp xong một tài liệu 43 trang rồi vẫn phải HỎI người
    dùng chép giúp bảng trong đó: chỉ có đường trích Fact (số kèm đơn vị) và trích bản đồ
    chân, còn mọi câu văn và bảng khác nằm ngoài tầm với."""
    agent = make_agent([])
    ctx = _ctx(agent)
    shutil.copy(tep_ban_giao, agent.config.paths.project_root / "bg.docx")
    agent.registry.run("doc.load", {"path": "bg.docx", "doc_id": "BG-3",
                                    "nguon": "noi_bo", "explain": _EX}, ctx)

    r = agent.registry.run("doc.read", {"doc_id": "BG-3", "tim": "Mức THẤP"}, ctx)
    assert r.ok, getattr(r.error, "message_vi", "")
    assert r.data["so_khop"] >= 1
    d = r.data["doan"][0]
    assert "4.2 Bảng chân đầy đủ" in d["trich_dan"] and d["la_bang"] is True

    # Lọc theo mục cũng phải dùng được, vì người hay nói "mục 4.2 viết gì".
    r2 = agent.registry.run("doc.read", {"doc_id": "BG-3", "muc": "4.2"}, ctx)
    assert r2.ok and r2.data["so_khop"] >= 3


def test_doc_read_khong_khop_thi_TU_CHOI_chu_khong_tra_ve_rong(make_agent, tep_ban_giao):
    """Trả một danh sách rỗng kèm `ok` là mời tác tử điền nốt bằng trí nhớ."""
    agent = make_agent([])
    ctx = _ctx(agent)
    shutil.copy(tep_ban_giao, agent.config.paths.project_root / "bg.docx")
    agent.registry.run("doc.load", {"path": "bg.docx", "doc_id": "BG-4",
                                    "nguon": "noi_bo", "explain": _EX}, ctx)
    r = agent.registry.run("doc.read", {"doc_id": "BG-4", "tim": "thuật toán PID"}, ctx)
    assert not r.ok and r.error.code == "E2004"
    assert "đừng lấy nội dung từ trí nhớ" in r.error.hint_for_agent


def test_doc_read_NOI_RA_khi_cat_bot_doan(make_agent, tep_ban_giao):
    agent = make_agent([])
    ctx = _ctx(agent)
    shutil.copy(tep_ban_giao, agent.config.paths.project_root / "bg.docx")
    agent.registry.run("doc.load", {"path": "bg.docx", "doc_id": "BG-5",
                                    "nguon": "noi_bo", "explain": _EX}, ctx)
    r = agent.registry.run("doc.read", {"doc_id": "BG-5", "gioi_han": 1}, ctx)
    assert r.ok and r.data["bi_cat"] is True
    assert "CÒN" in r.data["note_vi"] and "đoạn nữa" in r.data["note_vi"]


def test_fact_from_doc_KIEM_gia_tri_co_that_trong_doan(make_agent, tep_ban_giao):
    """Mô hình chọn đoạn và đặt tên; MÃ kiểm giá trị. Không có cửa nào để một con số nhớ
    được lọt vào kho qua đây."""
    agent = make_agent([])
    ctx = _ctx(agent)
    shutil.copy(tep_ban_giao, agent.config.paths.project_root / "bg.docx")
    agent.registry.run("doc.load", {"path": "bg.docx", "doc_id": "BG-6",
                                    "nguon": "noi_bo", "explain": _EX}, ctx)
    r = agent.registry.run("doc.read", {"doc_id": "BG-6", "tim": "400 kHz"}, ctx)
    so = r.data["doan"][0]["so"]

    ok = agent.registry.run("fact.from_doc", {
        "doc_id": "BG-6", "don_vi": so, "thuc_the": "bus:TWI", "khoa": "toc_do",
        "gia_tri": "400", "don_vi_do": "kHz"}, ctx)
    assert ok.ok, getattr(ok.error, "message_vi", "")
    assert "4.2 Bảng chân đầy đủ" in ok.data["trich_dan"]
    f = [x for x in agent.store.query_facts(subject="bus:TWI", limit=5)]
    assert f and f[0]["value"] == "400"
    src = f[0]["source"]
    if isinstance(src, str):
        import json as _j
        src = _j.loads(src)
    assert src["cite"] == ok.data["trich_dan"]

    xau = agent.registry.run("fact.from_doc", {
        "doc_id": "BG-6", "don_vi": so, "thuc_the": "bus:TWI", "khoa": "toc_do",
        "gia_tri": "1000", "don_vi_do": "kHz"}, ctx)
    assert not xau.ok and xau.error.code == "E2006"
    assert "KHÔNG có trong đoạn" in xau.error.message_vi


def test_fact_from_doc_nhan_ca_hai_dang_cua_mot_gia_tri(make_agent, tmp_path):
    """Tài liệu hay viết "39 (0x27)". Cả hai dạng đều là ĐỌC từ đoạn đó."""
    import docx

    agent = make_agent([])
    ctx = _ctx(agent)
    d = docx.Document()
    d.add_heading("12.5 Thanh ghi ba bộ định thời", level=2)
    d.add_paragraph("OCR2A = 39 (0x27) cho chu kỳ ngắt 20 µs.")
    p = agent.config.paths.project_root / "tg.docx"
    d.save(str(p))
    agent.registry.run("doc.load", {"path": "tg.docx", "doc_id": "TG-1",
                                    "nguon": "noi_bo", "explain": _EX}, ctx)
    r = agent.registry.run("doc.read", {"doc_id": "TG-1", "tim": "OCR2A"}, ctx)
    so = r.data["doan"][0]["so"]
    for gt in ("39", "0x27"):
        k = agent.registry.run("fact.from_doc", {
            "doc_id": "TG-1", "don_vi": so, "thuc_the": "reg:OCR2A", "khoa": "gia_tri",
            "gia_tri": gt}, ctx)
        assert k.ok, (gt, getattr(k.error, "message_vi", ""))
