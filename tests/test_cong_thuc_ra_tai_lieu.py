# -*- coding: utf-8 -*-
"""Công thức LaTeX phải thành ký hiệu người đọc được ở MỌI chỗ trong tài liệu xuất ra.

Anh Công báo: *"Bản xuất tài liệu của Agent đang lỗi không render latex ra file docx, pdf và
pptx"*. Đo lại trên một tài liệu có công thức ở sáu ngữ cảnh:

| Ngữ cảnh | Trước |
|---|---|
| Đoạn văn nguyên vẹn `$$…$$` | đổi được — **chỉ trong Word** |
| Giữa dòng `$…$` | in nguyên `\\frac{f_{VCO}}{PLLP}` |
| Gạch đầu dòng | in nguyên |
| Ô bảng | in nguyên |
| Trích dẫn | in nguyên |
| Rào ```math | in nguyên, dưới dạng khối mã |

Gốc: `cong_thuc_nguoi_doc` được gọi ở **đúng một chỗ** — nhánh đoạn văn của bộ dựng Word, và
chỉ khi cả đoạn là `$$…$$`. PowerPoint in cả dấu đô-la ra slide.

Bản vá đặt phép đổi vào **bộ quét chữ trong dòng** (`_quet`), nơi mọi ngữ cảnh đều đi qua.
"""

from __future__ import annotations

import pytest

from eide.xuat_ban import (chu_tran, cong_thuc_nguoi_doc, doc_markdown, render, tach_chu,
                           tex_con_sot)

NGUON = """# Thử công thức

$$f_{VCO} = f_{in} \\times \\frac{PLLN}{PLLM}$$

Tần số ra là $f_{out} = \\frac{f_{VCO}}{PLLP}$ với $PLLP \\in \\{2,4,6,8\\}$.

- Chu kỳ: $T = \\frac{1}{1000}$ s
- Sai số $\\pm 2\\%$

| Tham số | Công thức |
|---|---|
| VCO | $f_{in} \\times \\frac{N}{M}$ |

> Điện áp rơi $V_{drop} = I \\times R_{DS(on)}$ nhỏ hơn 100 mV.

```math
\\tau = R \\cdot C
```
"""


# ============================================================ mọi ngữ cảnh, không sót chỗ nào
@pytest.mark.parametrize("dinh_dang", ["docx", "pptx"])
def test_khong_con_cu_phap_TeX_nao_trong_tep(tmp_path, dinh_dang):
    """Phép kiểm chính: mở lại tệp và tìm dấu vết TeX. Còn một dấu là hỏng."""
    p = tmp_path / f"ra.{dinh_dang}"
    assert render(NGUON, p, dinh_dang).dat
    chu = _doc_lai(p, dinh_dang)
    for dau in ("\\frac", "\\times", "\\in", "\\pm", "\\cdot", "\\tau", "$$", "\\{"):
        assert dau not in chu, f"{dinh_dang} còn `{dau}` trên giấy:\n{chu[:400]}"


@pytest.mark.parametrize("dinh_dang", ["docx", "pptx"])
def test_ky_hieu_da_doi_CO_MAT_that(tmp_path, dinh_dang):
    """Chống rỗng: không có TeX vì đã đổi, chứ không phải vì công thức bị nuốt mất."""
    p = tmp_path / f"ra.{dinh_dang}"
    assert render(NGUON, p, dinh_dang).dat
    chu = _doc_lai(p, dinh_dang)
    for ky in ("×", "∈", "±", "·", "τ", "PLLN/PLLM", "f_VCO"):
        assert ky in chu, f"{dinh_dang} thiếu `{ky}`"


def _doc_lai(p, dinh_dang: str) -> str:
    if dinh_dang == "docx":
        from docx import Document

        d = Document(str(p))
        phan = [x.text for x in d.paragraphs]
        phan += [o.text for t in d.tables for h in t.rows for o in h.cells]
        return "\n".join(phan)
    from pptx import Presentation

    return "\n".join(
        "".join(r.text for r in par.runs)
        for s in Presentation(str(p)).slides for sh in s.shapes
        if sh.has_text_frame for par in sh.text_frame.paragraphs)


@pytest.mark.parametrize("ngu_canh,mau", [
    ("giữa dòng", "Tần số $f = \\frac{a}{b}$ Hz"),
    ("gạch đầu dòng", "- Chu kỳ $T = \\frac{1}{f}$ s"),
    ("trích dẫn", "> Rơi áp $V = I \\times R$ nhỏ"),
    ("tiêu đề", "## Công thức $E = mc^2$"),
])
def test_tung_ngu_canh_deu_doi_duoc(tmp_path, ngu_canh, mau):
    """Mỗi ngữ cảnh một ca riêng — hỏng một chỗ thì biết ngay là chỗ nào."""
    p = tmp_path / "ra.docx"
    assert render(f"# T\n\n{mau}\n", p, "docx").dat
    chu = _doc_lai(p, "docx")
    assert "\\" not in chu, f"{ngu_canh}: còn cú pháp TeX — {chu!r}"


def test_o_BANG_cung_doi_duoc(tmp_path):
    p = tmp_path / "ra.docx"
    assert render("# T\n\n| A | B |\n|---|---|\n| x | $\\frac{1}{2}$ |\n", p, "docx").dat
    assert "1/2" in _doc_lai(p, "docx")


def test_khoi_rao_math_KHONG_con_la_khoi_ma():
    """```math là CÔNG THỨC, không phải mã để đọc — cùng lý do với ```mermaid."""
    khoi = doc_markdown("# T\n\n```math\n\\tau = R \\cdot C\n```\n")
    assert [k.loai for k in khoi if k.loai in ("ma", "cong_thuc")] == ["cong_thuc"]


def test_doan_chi_co_cong_thuc_thanh_khoi_RIENG():
    """Nhận ra ở bộ đọc Markdown, một lần — để Word và PowerPoint cùng dùng.

    Trước đây phép nhận này nằm trong bộ dựng Word, nên PowerPoint in cả dấu đô-la ra slide.
    """
    khoi = doc_markdown("# T\n\n$$a = b$$\n")
    assert any(k.loai == "cong_thuc" and k.chu == "a = b" for k in khoi)


# ============================================================== dấu đô-la cũng là TIỀN
@pytest.mark.parametrize("cau", [
    "Giá $5 và $10 nữa, tổng $15.",
    "Chi phí ≈ 4 450 USD (khoảng $4,450).",
    "Mỗi lượt tốn $0,15 / 1M token đầu vào.",
    "Dùng shell: echo $HOME và $PATH",
    "Không đóng: giá $5 thôi",
    "Ba mức: $82 triệu · $111 triệu · $142 triệu",
])
def test_tien_KHONG_bi_doc_thanh_cong_thuc(cau):
    """Đây là chỗ nguy nhất của việc đọc `$…$` giữa dòng: ăn mất cả đoạn chữ ở giữa."""
    assert chu_tran(cau) == cau


@pytest.mark.parametrize("cau,mong", [
    ("Biến $x$ chạy", "Biến x chạy"),
    ("Công thức $E = mc^2$ đây", "Công thức E = mc^2 đây"),
    ("Tần số $f_{s}$ Hz", "Tần số f_s Hz"),
])
def test_cong_thuc_that_VAN_doi_duoc(cau, mong):
    """Chống vá quá tay: chốt chặn tiền không được giết luôn công thức thật."""
    assert chu_tran(cau) == mong


def test_so_thuan_trong_do_la_la_TIEN_khong_phai_cong_thuc():
    assert chu_tran("Giá $5$ thôi") == "Giá $5$ thôi"
    assert chu_tran("Giá $1.000$ đồng") == "Giá $1.000$ đồng"


# ====================================================== lệnh dài phải đổi trước lệnh ngắn
def test_MOI_lenh_trong_bang_doi_dung_mot_minh_no():
    """Phép kiểm phủ CẢ BẢNG — nó mới là thứ bắt được lỗi thứ tự.

    Ba bản trước đều đúng nhờ ba cơ chế khác nhau (thứ tự khai · `sorted` · lookahead), và
    **phá cái nào bộ kiểm cũng không đỏ** vì hai cái còn lại đỡ. Một ca chỉ thử `\\leq` không
    bắt được điều đó; ca này thử mọi lệnh, nên mọi cặp tiền tố đều bị soi.
    """
    from eide.xuat_ban import _TEX_KY_HIEU
    for lenh, ky in _TEX_KY_HIEU.items():
        ra = cong_thuc_nguoi_doc(lenh)
        if not ky.strip():
            # Lệnh BỐ CỤC (`\left`, `\quad`, `\,`…) đổi ra khoảng trắng hoặc rỗng. Một công
            # thức chỉ gồm một lệnh như thế thì ra chuỗi rỗng, vì hàm kết thúc bằng
            # `re.sub(r"\s+", " ")` rồi `.strip()`. Đó là **đúng**: trong một dòng chữ thuần,
            # lệnh bố cục không có gì để hiện.
            #
            # Ca này từng so `ra == ky` cho mọi mục, nên nó đỏ ngay lúc tôi thêm mười ba lệnh
            # bố cục vào bảng ngày 02/10/2026 — và nó đỏ ĐÚNG, vì bảng khi ấy hứa `\quad` ra
            # hai dấu cách trong khi bước chuẩn hoá sau đó xoá đi. Lời hứa đã sửa; chỗ cần
            # nói rõ là ca kiểm hỏi gì với nhóm mục này.
            # Và phải thử TRONG NGỮ CẢNH, không thử đứng một mình: hàm mở đầu bằng
            # `s.strip()`, nên `\ ` (gạch chéo + dấu cách) mất dấu cách trước khi bảng chạy
            # và còn trơ lại một gạch chéo. Đó là chuyện của phép thử, không phải của bảng —
            # trong một công thức thật thì lệnh ấy luôn nằm giữa hai thứ khác.
            #
            # Ngữ cảnh phải chọn theo DẠNG TÊN LỆNH. Lệnh viết bằng chữ cần một dấu không
            # phải chữ ngăn phía sau, vì `\leftb` KHÔNG phải `\left` — và chốt
            # `(?![A-Za-z])` từ chối đúng. Lệnh dạng dấu câu thì tên kết thúc ở dấu ấy, nên
            # `a\,b` là cách dùng thường và phải đổi được.
            sau = "(" if lenh[1:2].isalpha() else "b"
            trong = cong_thuc_nguoi_doc(f"a{lenh}{sau}")
            assert "\\" not in trong, (
                f"{lenh!r} là lệnh bố cục mà còn để lại gạch chéo: a{lenh}{sau} → {trong!r}")
            assert trong.replace(" ", "") == f"a{sau}", (
                f"{lenh!r} phải biến mất: a{lenh}{sau} → {trong!r}")
            continue
        assert ra == ky, f"{lenh} → {ra!r}, đợi {ky!r}"


def test_cap_tien_to_khong_an_mat_nhau():
    """`\\le` là tiền tố của `\\leq`; đổi ngắn trước thì để lại một chữ `q` lạc giữa công thức."""
    from eide.xuat_ban import _TEX_KY_HIEU
    cap = [(a, b) for a in _TEX_KY_HIEU for b in _TEX_KY_HIEU
           if a != b and b.startswith(a)]
    assert cap, "bảng không còn cặp tiền tố nào — ca này hết tác dụng, xem lại"
    for ngan, dai in cap:
        assert cong_thuc_nguoi_doc(dai) == _TEX_KY_HIEU[dai], \
            f"`{dai}` bị `{ngan}` ăn mất"


def test_ky_tu_thoat_thanh_chinh_no():
    assert cong_thuc_nguoi_doc(r"2\%") == "2%"
    assert cong_thuc_nguoi_doc(r"\{a,b\}") == "{a,b}"


def test_du_bo_chu_Hy_Lap():
    """Thiếu một chữ là một công thức in ra cú pháp thô — nên đủ bộ, không chọn lọc."""
    for lenh, ky in ((r"\tau", "τ"), (r"\sigma", "σ"), (r"\lambda", "λ"),
                     (r"\theta", "θ"), (r"\phi", "φ"), (r"\rho", "ρ")):
        assert cong_thuc_nguoi_doc(lenh) == ky


# ================================================ nói ra thứ KHÔNG đổi được, đừng im lặng
def test_lenh_la_duoc_KE_RA_chu_khong_im():
    khoi = doc_markdown("# T\n\nCó $\\varrho_{xy}$ ở đây.\n")
    assert tex_con_sot(khoi) == ["\\varrho"]


def test_lenh_DA_doi_duoc_thi_khong_bao_dong():
    """Soi bản ĐÃ đổi, không soi nguồn: một cảnh báo luôn kêu thì bằng không kêu."""
    khoi = doc_markdown("# T\n\n$a \\times \\frac{b}{c} \\leq d$\n")
    assert tex_con_sot(khoi) == []


def test_khoi_MA_khong_bi_tinh_la_cong_thuc_sot():
    """Mã giữ nguyên văn là đúng — đếm nó vào phần sót sẽ báo động mỗi lần có LaTeX trong ví dụ."""
    assert tex_con_sot(doc_markdown("# T\n\n```c\nint x; // \\n newline\n```\n")) == []


def test_render_bao_phan_chua_doi_duoc(tmp_path):
    kq = render("# T\n\nCó $\\varrho$ ở đây.\n", tmp_path / "r.docx", "docx")
    assert kq.dat and kq.do_lai.get("tex_chua_doi_duoc") == ["\\varrho"]


def test_render_sach_thi_KHONG_co_truong_canh_bao(tmp_path):
    kq = render("# T\n\nCó $a \\times b$ ở đây.\n", tmp_path / "r.docx", "docx")
    assert "tex_chua_doi_duoc" not in kq.do_lai


# ============================================================ công thức không phá kiểu chữ
def test_cong_thuc_doc_TRUOC_dau_nhan_manh():
    """`$a * b$` có dấu sao; để nhánh nghiêng đọc trước thì công thức bị cắt làm đôi."""
    ra = tach_chu("Ta có $a * b = c$ nhé")
    cong = [c for c, k in ra if "cong_thuc" in k]
    assert cong == ["a * b = c"]


def test_chu_dam_ngoai_cong_thuc_van_dam():
    ra = tach_chu("**Quan trọng: $x = 1$**")
    assert all("dam" in k for _, k in ra)
    assert any("cong_thuc" in k for _, k in ra)


def test_phan_so_chi_dong_ngoac_khi_CAN():
    """`(1)/(1000)` đúng nhưng đọc vướng, và trên một trang đầy công thức thì vướng cộng dồn.

    Ngoặc chỉ cần khi bỏ đi sẽ đổi nghĩa — tức khi vế có phép toán hoặc khoảng trắng.
    """
    assert cong_thuc_nguoi_doc(r"\frac{1}{1000}") == "1/1000"
    assert cong_thuc_nguoi_doc(r"\frac{PLLN}{PLLM}") == "PLLN/PLLM"
    assert cong_thuc_nguoi_doc(r"\frac{a+b}{c}") == "(a+b)/c"
    assert cong_thuc_nguoi_doc(r"\frac{a \times b}{c}") == "(a × b)/c"


# ======================================= ba lỗi chỉ trang PDF bắt được, ngày 01/10/2026
def test_can_bac_hai_go_ngoac_nhon():
    """`\\sqrt` nằm trong bảng ký hiệu như một ký tự lẻ, nên `\\sqrt{R^2+X^2}` ra `√{R^2+X^2}`
    — ngoặc nhọn của TeX lọt nguyên ra giấy."""
    assert cong_thuc_nguoi_doc(r"\sqrt{R^2 + X^2}") == "√(R^2 + X^2)"
    assert cong_thuc_nguoi_doc(r"\sqrt{2}") == "√2"


def test_do_la_HAU_TO_khong_phai_so_mu():
    """`2^{\\circ}` ra `2^°` vì luật `^{…}` chạy trước — mà độ không phải số mũ."""
    assert cong_thuc_nguoi_doc(r"\pm 2^{\circ}") == "± 2°"
    assert cong_thuc_nguoi_doc(r"45^\circ") == "45°"


def test_mau_nhieu_HANG_phai_dong_ngoac():
    """`1/2πτ` **đọc thành `(1/2)·π·τ`** — sai nghĩa, không chỉ xấu.

    Đếm theo ký tự phép toán thì không bắt được: `2\\pi\\tau` không có dấu cộng hay khoảng
    trắng nào. Phải đếm theo HẠNG — một lệnh TeX, một tên, hay một số là một hạng.
    """
    assert cong_thuc_nguoi_doc(r"\frac{1}{2\pi\tau}") == "1/(2πτ)"
    assert cong_thuc_nguoi_doc(r"\frac{1}{2R}") == "1/(2R)"
    # và không đóng ngoặc khi chỉ có một hạng
    assert cong_thuc_nguoi_doc(r"\frac{1}{1000}") == "1/1000"
    assert cong_thuc_nguoi_doc(r"\frac{f_{VCO}}{PLLP}") == "f_VCO/PLLP"
